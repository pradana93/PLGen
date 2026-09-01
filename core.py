import tkinter as tk
from tkinter import ttk, messagebox
import datetime
import os
import json
import sys
import requests
import threading
import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage
import addons

# --- NEW: Vault Encryption Library ---
from cryptography.fernet import Fernet
# --------------------------------

# --- NEW: Drag & Drop Library ---
try:
    from tkinterdnd2 import DND_FILES, TkinterDnD
    DND_AVAILABLE = True
except ImportError:
    DND_AVAILABLE = False
# --------------------------------

# --- NEW: Audio Feedback System ---
try:
    import winsound
except ImportError:
    winsound = None

def beep_success():
    if winsound: winsound.MessageBeep(winsound.MB_OK)

def beep_error():
    if winsound: winsound.MessageBeep(winsound.MB_ICONHAND)
# ----------------------------------


# ==========================================
#     DYNAMIC CLOUD MASTER DATA ENGINE
# ==========================================

# The Ultimate Doomsday Fallbacks (In case Cloud & Local Cache both fail)
FALLBACK_CHECKERS = ["Aji", "Luthfi", "Fadly", "Masroor"]

FALLBACK_MASTER_DATA = {
    "BOX_TOLERANCE": 1.859,
    "CATEGORIES": {
        "FROZEN_ITEMS": ["Beef Patty Small", "Beef Patty Large", "Keju Slice Non Brand", "Keju Slice Non Brand A", "Chicken Nugget", "Spicy Chicken Nugget", "Bangor Fried Chicken", "Sosis", "Ayam Crispy", "Dori Crispy", "BEEF SLICE", "Smoke Beef Slice", "Spicy Chicken Patty", "Bangor Chicken Wings"],
        "KENTANG_ITEMS": ["Kentang Goreng", "Kentang Goreng Mc Cain"],
        "SAUCE_ITEMS": ["BBQ Sauce", "BBQ Spicy", "Bolognese Sauce 500gr", "Cheese Sauce", "Nacho Sauce New", "Mayonaise Garlic", "Nestea Lemontea", "Butter", "Thousand Island Mayonaise"],
        "PACKAGING_ITEMS": ["Kertas Nasi", "Paper Kentang", "Tray Kentang", "Paper Bag", "Packaging Box Sultan", "Packaging HD", "Kertas Printer", "Cup Plastik 14 oz", "Tutup Gelas", "Cup Sauce", "Sedotan", "Hand Gloves", "Kresek Kecil", "Kresek Besar", "Kresek Gelas", "Bangor Crazy Bucket", "Spunbond Bangor", "Sticker Labeling", "Tissue Pop Up", "Box Hampers", "Grill Box", "Inner", "Bangor Thermal Bag", "Kertas Thermal", "Kupon Umroh"],
        "APPAREL_ITEMS": ["Kaus Seragam M", "Kaus Seragam L", "Kaus Seragam XL", "Kaus Seragam XXL", "Kaus Seragam XXXL", "Topi", "Polo Shirt S", "Polo Shirt M", "Polo Shirt L", "Polo Shirt XL", "Polo Shirt XXL", "Polo Shirt XXXL", "Apron", "Seragam Owner S", "Seragam Owner M", "Seragam Owner L", "Seragam Owner XL", "Seragam Owner XXL", "Seragam Owner XXXL"],
        "BIG_ITEMS": ["Minyak Padat", "Sabun Cuci Piring Mitra", "Sabun Cuci Tangan @5 Liter", "Sabun Lantai Mitra", "Hand Sanitizer", "Sabun MPC", "Sabun Kerak", "Cetakan Telur", "Sambal Sachet Bangor", "Saos Tomat Jerigen Delmonte"],
        "BREAD_ITEMS": ["HD Bun", "Burger Bun"],
        "BUNDLE_ITEMS": ["Box Hampers", "Grill Box", "Inner"]
    },
    "ITEM_UOM": {
        "Beef Patty Small": "Pack", "Beef Patty Large": "Pack", "Keju Slice Non Brand": "Pack", "Keju Slice Non Brand A": "Pack", "Chicken Nugget": "Pack", "Spicy Chicken Nugget": "Pack", "Kentang Goreng": "Pack", "Kentang Goreng Mc Cain": "Pack", "Bangor Fried Chicken": "Pack", "Sosis": "Pack", "Ayam Crispy": "Pack", "Dori Crispy": "Pack", "BEEF SLICE": "Pack", "Smoke Beef Slice": "Pack", "Spicy Chicken Patty": "Pack", "HD Bun": "Pack", "Burger Bun": "Pack", "Thousand Island Mayonaise": "Pack", "BBQ Sauce": "Pack", "BBQ Spicy": "Pack", "Bolognese Sauce 500gr": "Pack", "Cheese Sauce": "Pack", "Nacho Sauce New": "Pack", "Mayonaise Garlic": "Pack", "Nestea Lemontea": "Pack", "Sambal Sachet Bangor": "Dus", "Butter": "Pack", "Sticker Labeling": "Ikat", "Tissue Pop Up": "Pack", "Kertas Nasi": "Pack", "Paper Kentang": "Pack", "Tray Kentang": "Pack", "Paper Bag": "Pack", "Packaging HD": "Ikat", "Cup Plastik 14 oz": "Pack", "Tutup Gelas": "Pack", "Cup Sauce": "Pack", "Sedotan": "Pack", "Hand Gloves": "Pack", "Kresek Kecil": "Pack", "Kresek Besar": "Pack", "Kresek Gelas": "Pack", "": "Pack", "Bangor Thermal Bag": "Pack", "Kupon Umroh": "Buku", "Packaging Box Sultan": "Ikat", "Minyak Padat": "Dus", "Box Hampers": "Pcs", "Grill Box": "Pcs", "Inner": "Pcs", "Sabun Cuci Piring Mitra": "Jrg", "Sabun Cuci Tangan @5 Liter": "Jrg", "Sabun Lantai Mitra": "Jrg", "Hand Sanitizer": "Jrg", "Sabun MPC": "Jrg", "Sabun Kerak": "Jrg", "Kaus Seragam M": "Pcs", "Kaus Seragam L": "Pcs", "Kaus Seragam XL": "Pcs", "Kaus Seragam XXL": "Pcs", "Kaus Seragam XXXL": "Pcs", "Topi": "Pcs", "Bangor Crazy Bucket": "Pack", "Polo Shirt S": "Pcs", "Polo Shirt M": "Pcs", "Polo Shirt L": "Pcs", "Polo Shirt XL": "Pcs", "Polo Shirt XXL": "Pcs", "Polo Shirt XXXL": "Pcs", "Apron": "Pcs", "Seragam Owner S": "Pcs", "Seragam Owner M": "Pcs", "Seragam Owner L": "Pcs", "Seragam Owner XL": "Pcs", "Seragam Owner XXL": "Pcs", "Seragam Owner XXXL": "Pcs", "Cetakan Telur": "Pcs", "Bangor Chicken Wings": "Pack", "Kertas Thermal": "Pack", "Saos Tomat Jerigen Delmonte": "Jrg"
    },
    "BOX_CAPACITY": {
        "Beef Patty Small": 18, "Beef Patty Large": 18, "Keju Slice Non Brand": 12, "Keju Slice Non Brand A": 12, "Chicken Nugget": 10, "HD Bun": 30, "Burger Bun": 20, "Kentang Goreng": 15, "Kentang Goreng Mc Cain": 20, "Bangor Fried Chicken": 6, "Spicy Chicken Nugget": 5, "Spicy Chicken Patty": 10, "Sosis": 10, "Ayam Crispy": 10, "Dori Crispy": 24, "BEEF SLICE": 40, "Smoke Beef Slice": 150, "Thousand Island Mayonaise": 20, "BBQ Sauce": 20, "BBQ Spicy": 20, "Bolognese Sauce 500gr": 20, "Cheese Sauce": 12, "Nacho Sauce New": 24, "Mayonaise Garlic": 20, "Nestea Lemontea": 12, "Butter": 40, "Sambal Sachet Bangor": 1, "Minyak Padat": 1, "Kertas Nasi": 20, "Paper Kentang": 30, "Tray Kentang": 60, "Paper Bag": 30, "Packaging Box Sultan": 50, "Packaging HD": 50, "Kertas Printer": 10, "Cup Plastik 14 oz": 40, "Tutup Gelas": 40, "Cup Sauce": 24, "Sedotan": 50, "Hand Gloves": 150, "Kresek Kecil": 50, "Kresek Besar": 50, "Kresek Gelas": 50, "Kaus Seragam M": 50, "Kaus Seragam L": 50, "Kaus Seragam XL": 50, "Kaus Seragam XXL": 50, "Kaus Seragam XXXL": 50, "Topi": 50, "Box Hampers": 50, "Grill Box": 50, "Inner": 50, "Bangor Crazy Bucket": 9, "Spunbond Bangor": 50, "Bangor Thermal Bag": 5, "Sticker Labeling": 100, "Polo Shirt S": 50, "Polo Shirt M": 50, "Polo Shirt L": 50, "Polo Shirt XL": 50, "Polo Shirt XXL": 50, "Polo Shirt XXXL": 50, "Apron": 50, "Seragam Owner S": 50, "Seragam Owner M": 50, "Seragam Owner L": 50, "Seragam Owner XL": 50, "Seragam Owner XXL": 50, "Seragam Owner XXXL": 50, "Cetakan Telur": 2, "Sabun Cuci Piring Mitra": 1, "Sabun Cuci Tangan @5 Liter": 1, "Sabun Lantai Mitra": 1, "Hand Sanitizer": 1, "Sabun MPC": 1, "Sabun Kerak": 1, "Tissue Pop Up": 50, "Bangor Chicken Wings": 6, "Kertas Thermal": 10, "Saos Tomat Jerigen Delmonte": 1, "Kupon Umroh": 200
    }
}

def sync_master_data():
    """Fetches Master Data from Cloud. Uses Local Cache if offline. Falls back to hardcode if catastrophic."""
    CLOUD_URL = "https://jestu93.pythonanywhere.com/static/master_data.json"
    
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))
        
    cache_dir = os.path.join(base_dir, "logs")
    os.makedirs(cache_dir, exist_ok=True)
    cache_file = os.path.join(cache_dir, "master_data_cache.json")
    
    try:
        # 1. Attempt Cloud Sync
        response = requests.get(CLOUD_URL, timeout=4)
        if response.status_code == 200:
            cloud_data = response.json()
            # Save Backup
            with open(cache_file, "w") as f:
                json.dump(cloud_data, f, indent=4)
            return cloud_data, "Cloud Sync ☁️"
    except Exception:
        pass # Offline or server down
        
    # 2. Attempt Local Cache
    if os.path.exists(cache_file):
        try:
            with open(cache_file, "r") as f:
                cache_data = json.load(f)
            return cache_data, "Local Backup 💾"
        except Exception:
            pass 
            
    # 3. Doomsday Fallback
    return FALLBACK_MASTER_DATA, "Hardcoded Backup ⚠️"

def fetch_checkers():
    """Fetches the dynamic list of checkers from the server. Falls back to hardcoded list if offline."""
    try:
        response = requests.get(f"{addons.BASE_URL}/api/checkers", timeout=3)
        if response.status_code == 200:
            data = response.json()
            checkers = data.get("checkers", [])
            if checkers:
                return checkers
    except Exception:
        pass
    return FALLBACK_CHECKERS

# Execute Data Load
master_db, DATA_SOURCE_STATE = sync_master_data()

# MAP TO GLOBAL VARIABLES
MASTER_BOX_TOLERANCE = master_db.get("BOX_TOLERANCE", 1.859)
FROZEN_ITEMS = set(master_db["CATEGORIES"].get("FROZEN_ITEMS", []))
KENTANG_ITEMS = set(master_db["CATEGORIES"].get("KENTANG_ITEMS", []))
SAUCE_ITEMS = set(master_db["CATEGORIES"].get("SAUCE_ITEMS", []))
PACKAGING_ITEMS = set(master_db["CATEGORIES"].get("PACKAGING_ITEMS", []))
APPAREL_ITEMS = set(master_db["CATEGORIES"].get("APPAREL_ITEMS", []))
BIG_ITEMS = set(master_db["CATEGORIES"].get("BIG_ITEMS", []))
BREAD_ITEMS = set(master_db["CATEGORIES"].get("BREAD_ITEMS", []))
BUNDLE_ITEMS = set(master_db["CATEGORIES"].get("BUNDLE_ITEMS", []))

# Combined DRY_ITEMS for the packing logic
DRY_ITEMS = (SAUCE_ITEMS | PACKAGING_ITEMS | APPAREL_ITEMS) - BUNDLE_ITEMS

ITEM_UOM = master_db.get("ITEM_UOM", {})
box_capacity = master_db.get("BOX_CAPACITY", {})
sku_lookup = {sku.lower(): sku for sku in box_capacity.keys()}


# ==========================================
#     CORE LOGIC
# ==========================================

def calculate_boxes(order, current_tolerance=1.859):
    final_boxes = []
    remaining = {sku: data["qty"] for sku, data in order.items() if data["qty"] > 0}
    TOLERANCE = current_tolerance

    isolated_keju_boxes = []
    for sku in ["Keju Slice Non Brand", "Keju Slice Non Brand A"]:
        if sku in remaining and order[sku]["qty"] > 3:
            cap = box_capacity.get(sku, 12)
            while remaining[sku] > 0:
                take = min(remaining[sku], cap)
                isolated_keju_boxes.append({sku: take})
                remaining[sku] -= take
            del remaining[sku]

    def pack_category_greedy(category_set):
        active_set = category_set - BUNDLE_ITEMS
        cat_skus = [s for s in active_set if s in remaining and remaining[s] > 0]
        cat_skus.sort(key=lambda x: 1/box_capacity.get(x, 1), reverse=True)
        
        current_box = {}
        used_space = 0.0

        for sku in cat_skus:
            qty_to_pack = remaining[sku]
            if qty_to_pack <= 0:
                continue

            unit_space = 1 / box_capacity.get(sku, 1)
            item_total_space = qty_to_pack * unit_space

            if current_box and (used_space + item_total_space > TOLERANCE + 0.0001):
                final_boxes.append(current_box)
                current_box = {}
                used_space = 0.0
                
            current_box[sku] = qty_to_pack
            used_space += item_total_space
            remaining[sku] = 0
        
        if current_box: 
            final_boxes.append(current_box)

    for sku in list(remaining.keys()):
        if sku in FROZEN_ITEMS:
            cap = box_capacity.get(sku, 1)
            if sku in ["Beef Patty Large", "Beef Patty Small"] or remaining[sku] >= cap:
                full_boxes = remaining[sku] // cap
                for _ in range(full_boxes): final_boxes.append({sku: cap})
                remaining[sku] %= cap
                if remaining[sku] == 0: del remaining[sku]

    for sku in KENTANG_ITEMS:
        if sku in remaining:
            cap = box_capacity.get(sku, 15)
            if remaining[sku] >= cap:
                full_boxes = remaining[sku] // cap
                for _ in range(full_boxes): final_boxes.append({sku: cap})
                remaining[sku] %= cap

    if isolated_keju_boxes:
        final_boxes.extend(isolated_keju_boxes)

    pack_category_greedy(FROZEN_ITEMS)
    pack_category_greedy(KENTANG_ITEMS)

    for sku in list(remaining.keys()):
        if sku in DRY_ITEMS:
            cap = box_capacity.get(sku, 1)
            if cap > 1 and remaining[sku] >= cap:
                full_boxes = remaining[sku] // cap
                for _ in range(full_boxes): final_boxes.append({sku: cap})
                remaining[sku] %= cap

    pack_category_greedy(DRY_ITEMS)

    for sku in BREAD_ITEMS:
        if sku in remaining:
            cap = box_capacity.get(sku, 1)
            while remaining[sku] > 0:
                take = min(remaining[sku], cap)
                final_boxes.append({sku: take})
                remaining[sku] -= take
    
    for sku in BIG_ITEMS:
        if sku in remaining:
            for _ in range(remaining[sku]): final_boxes.append({sku: 1})
            remaining[sku] = 0

    current_bundle_box = {}
    current_bundle_qty = 0
    for sku in BUNDLE_ITEMS:
        while sku in remaining and remaining[sku] > 0:
            space_left = 50 - current_bundle_qty
            take = min(remaining[sku], space_left)
            
            if take > 0:
                current_bundle_box[sku] = current_bundle_box.get(sku, 0) + take
                remaining[sku] -= take
                current_bundle_qty += take
            
            if current_bundle_qty == 50:
                final_boxes.append(current_bundle_box)
                current_bundle_box = {}
                current_bundle_qty = 0
                
        if sku in remaining:
            del remaining[sku]
            
    if current_bundle_box:
        final_boxes.append(current_bundle_box)

    return final_boxes

def get_next_delivery_number(notes_combined, company_code="BBB"):
    os.makedirs("logs", exist_ok=True)
    counter_file = os.path.join("logs", f"delivery_counter_{company_code}.txt")
    today = datetime.datetime.now().strftime("%d%m%Y")
    
    last_number, last_date = 0, None
    if os.path.exists(counter_file):
        with open(counter_file, "r") as f:
            content = f.read().strip().split("|")
            if len(content) == 2: last_date, last_number = content[0], int(content[1])
            
    if last_date != today: last_number = 0
    new_number = last_number + 1
    
    with open(counter_file, "w") as f: f.write(f"{today}|{new_number}")
    return f"DO/{company_code}/{today}/{new_number:03d}"

# ==========================================
#     NEW: OUTLET VALIDATION & AUTO-REGISTRATION
# ==========================================
def validate_or_register_outlet(parent, outlet_name):
    """
    Validates if outlet exists in Master Data.
    If not, shows Typo Guard dialog to either select a similar outlet or register new.
    Returns (True, final_outlet_name) if validated/registered, (False, None) if cancelled.
    """
    outlet_upper = outlet_name.strip().upper()
    
    # Check if already registered
    if outlet_upper in master_db.get("OUTLET_INFO", {}):
        return True, outlet_name
    
    # Find similar outlets (simple fuzzy matching - substring or contains)
    all_outlets = list(master_db.get("OUTLET_INFO", {}).keys())
    similar = [o for o in all_outlets if outlet_upper in o or o in outlet_upper]
    
    # Show Typo Guard Dialog
    dialog = tk.Toplevel(parent)
    dialog.title("🆕 New Outlet Detected")
    dialog.geometry("450x550")
    dialog.attributes("-topmost", True)
    dialog.grab_set()
    dialog.configure(bg="#f4f6f9")
    
    tk.Label(dialog, text="🆕 NEW OUTLET DETECTED", 
             font=(addons.PREMIUM_FONT, 14, "bold"), 
             bg="#f4f6f9", fg="#e74c3c").pack(pady=(20, 10))
    
    tk.Label(dialog, text=f'"{outlet_name}" is not in Master Data.', 
             font=(addons.PREMIUM_FONT, 10), 
             bg="#f4f6f9", fg="#2c3e50").pack(pady=(0, 10))
    
    # Similar outlets section
    selection_var = tk.StringVar(value="")
    
    if similar:
        tk.Label(dialog, text="Did you mean one of these?", 
                 font=(addons.PREMIUM_FONT, 10, "bold"), 
                 bg="#f4f6f9", fg="#2c3e50").pack(pady=(10, 5))
        
        for outlet in similar[:5]:  # Show max 5 suggestions
            rb = tk.Radiobutton(dialog, text=outlet, variable=selection_var, value=outlet,
                               font=(addons.PREMIUM_FONT, 10), bg="#f4f6f9", anchor="w")
            rb.pack(fill="x", padx=30)
        
        tk.Radiobutton(dialog, text="No — create it as a NEW outlet", 
                      variable=selection_var, value="__CREATE_NEW__",
                      font=(addons.PREMIUM_FONT, 10, "bold"), 
                      bg="#f4f6f9", fg="#e74c3c", anchor="w").pack(fill="x", padx=30, pady=(10, 0))
    else:
        selection_var.set("__CREATE_NEW__")
        tk.Label(dialog, text="No similar outlets found. Creating as new outlet.", 
                 font=(addons.PREMIUM_FONT, 10, "italic"), 
                 bg="#f4f6f9", fg="#7f8c8d").pack(pady=10)
    
    # Optional fields
    details_frame = tk.Frame(dialog, bg="#ecf0f1", bd=1, relief="solid")
    details_frame.pack(fill="x", padx=20, pady=15)
    
    tk.Label(details_frame, text="Optional Details (Admin can complete later):", 
             font=(addons.PREMIUM_FONT, 9, "bold"), 
             bg="#ecf0f1", fg="#2c3e50").grid(row=0, column=0, columnspan=2, sticky="w", padx=10, pady=(8, 5))
    
    tk.Label(details_frame, text="Receiver:", font=(addons.PREMIUM_FONT, 9), bg="#ecf0f1").grid(row=1, column=0, sticky="w", padx=10, pady=2)
    receiver_var = tk.StringVar()
    tk.Entry(details_frame, textvariable=receiver_var, font=(addons.PREMIUM_FONT, 9), width=30).grid(row=1, column=1, padx=10, pady=2)
    
    tk.Label(details_frame, text="Phone:", font=(addons.PREMIUM_FONT, 9), bg="#ecf0f1").grid(row=2, column=0, sticky="w", padx=10, pady=2)
    phone_var = tk.StringVar()
    tk.Entry(details_frame, textvariable=phone_var, font=(addons.PREMIUM_FONT, 9), width=30).grid(row=2, column=1, padx=10, pady=2)
    
    tk.Label(details_frame, text="Address:", font=(addons.PREMIUM_FONT, 9), bg="#ecf0f1").grid(row=3, column=0, sticky="w", padx=10, pady=2)
    address_var = tk.StringVar()
    tk.Entry(details_frame, textvariable=address_var, font=(addons.PREMIUM_FONT, 9), width=30).grid(row=3, column=1, padx=10, pady=2)
    
    result = {"status": False, "final_outlet": None}
    
    def on_confirm():
        selected = selection_var.get()
        
        if selected == "__CREATE_NEW__":
            # Register new outlet
            new_outlet = {
                "name": outlet_name,
                "phone": phone_var.get().strip(),
                "address": address_var.get().strip(),
                "auto_registered": True,
                "registered_by": addons.HWID_ALIASES.get(addons.get_hwid(), "Unknown"),
                "registered_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            
            # Add to local master_db
            if "OUTLET_INFO" not in master_db:
                master_db["OUTLET_INFO"] = {}
            master_db["OUTLET_INFO"][outlet_upper] = new_outlet
            
            # Save to cache
            try:
                cache_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "logs")
                os.makedirs(cache_dir, exist_ok=True)
                cache_file = os.path.join(cache_dir, "master_data_cache.json")
                with open(cache_file, "w") as f:
                    json.dump(master_db, f, indent=4)
            except Exception as e:
                print(f"Failed to save cache: {e}")
            
            # Push to server (or queue if offline)
            def push_to_server():
                try:
                    payload = {
                        "outlet_name": outlet_upper,
                        "outlet_data": new_outlet
                    }
                    requests.post(f"{addons.BASE_URL}/api/outlet/register", 
                                json=payload, headers=addons.API_HEADERS, timeout=5)
                except Exception:
                    # Queue for later sync
                    try:
                        queue_file = os.path.join("logs", "pending_sync.json")
                        queue = []
                        if os.path.exists(queue_file):
                            with open(queue_file, "r") as f:
                                queue = json.load(f)
                        queue.append({
                            "endpoint": "/api/outlet/register",
                            "payload": payload
                        })
                        with open(queue_file, "w") as f:
                            json.dump(queue, f)
                    except:
                        pass
            
            threading.Thread(target=push_to_server, daemon=True).start()
            
            # Audit log
            addons.log_action(f"🆕 OUTLET AUTO-REGISTERED: {outlet_upper} by {new_outlet['registered_by']}")
            
            result["status"] = True
            result["final_outlet"] = outlet_name
            dialog.destroy()
            
        elif selected and selected != "__CREATE_NEW__":
            # User selected an existing outlet
            result["status"] = True
            result["final_outlet"] = selected
            dialog.destroy()
        else:
            messagebox.showwarning("Selection Required", "Please select an outlet or create a new one.", parent=dialog)
    
    def on_cancel():
        result["status"] = False
        dialog.destroy()
    
    btn_frame = tk.Frame(dialog, bg="#f4f6f9")
    btn_frame.pack(fill="x", padx=20, pady=15)
    
    tk.Button(btn_frame, text="✅ CONFIRM", command=on_confirm,
             bg="#27ae60", fg="white", font=(addons.PREMIUM_FONT, 11, "bold"),
             width=15).pack(side="left", expand=True, padx=5)
    
    tk.Button(btn_frame, text="❌ CANCEL", command=on_cancel,
             bg="#95a5a6", fg="white", font=(addons.PREMIUM_FONT, 11, "bold"),
             width=15).pack(side="right", expand=True, padx=5)
    
    dialog.protocol("WM_DELETE_WINDOW", on_cancel)
    parent.wait_window(dialog)
    
    return result["status"], result.get("final_outlet")
# ==========================================

def export_packing_list(outlet_name, boxes, order):
    os.makedirs("logs", exist_ok=True)
    os.makedirs("Labels", exist_ok=True) 
    
    delivery_date = addons.get_delivery_date()
    
    today = datetime.datetime.now()
    pl_filename = f"{today.strftime('%d%m%Y')}_{outlet_name}.xlsx"
    pl_filepath = os.path.join("logs", pl_filename)
    
    lbl_filename = f"{outlet_name}_{today.strftime('%d%m%Y')}.xlsx"
    lbl_filepath = os.path.join("Labels", lbl_filename)
    
    template_path = os.path.join("logs", "template.xlsx")
    
    unique_notes = sorted(set(item["note"] for item in order.values() if item.get("note")))
    notes_combined = " & ".join(unique_notes) if unique_notes else "N/A"
    
    delivery_no = get_next_delivery_number(notes_combined, addons.ACTIVE_COMPANY_CODE)

    # --- NEW: Generate QR Code for Tracking ---
    from urllib.parse import quote
    encoded_delivery_no = quote(delivery_no, safe='')  # Encode slashes as %2F
    qr_url = f"{addons.BASE_URL}/scan/{encoded_delivery_no}"
    
    # FIX: Sanitize the delivery_no to prevent directory/folder errors
    safe_delivery_no = delivery_no.replace("/", "_")
    qr_img_path = os.path.join("logs", f"qr_{safe_delivery_no}.png")
    
    addons.generate_qr_code_image(qr_url, qr_img_path)
    # ------------------------------------------

    addons.log_to_shift_report(order)

    # --- NEW: Calculate Total Weight for Outbound Manifest ---
    item_weights = master_db.get("ITEM_WEIGHT_GRAMS", {})
    total_weight_kg = 0.0
    for sku, data in order.items():
        qty = data["qty"]
        weight_g = item_weights.get(sku, 0)
        total_weight_kg += (qty * weight_g) / 1000.0
    # --------------------------------------------------------

    # --- NEW: Grab Checker and Cluster Display Text ---
    selected_checker = checker_var.get()
    cluster_text = cluster_var.get().strip()

    if selected_checker == "Select Checker" or not selected_checker:
        messagebox.showwarning("Missing Info", "Please select a Checker before exporting.")
        return

    checker_display = f"{selected_checker} | Cluster: {cluster_text}" if cluster_text else selected_checker
    # ------------------------------------------------

    display_rows = []
    for i, box in enumerate(boxes, start=1):
        is_first_in_koli = True
        for sku, qty in box.items():
            display_rows.append({
                "koli": i if is_first_in_koli else "",
                "sku": sku,
                "qty": qty,
                "uom": ITEM_UOM.get(sku, "Pack"),
                "note": order.get(sku, {}).get("note", "")
            })
            is_first_in_koli = False

    MAX_ROWS_PER_PAGE = 34 
    START_ROW = 9

    if os.path.exists(template_path):
        wb = load_workbook(template_path)
        master_sheet = wb.active
        master_sheet.title = "TEMP_MASTER" 
        num_pages = (len(display_rows) + MAX_ROWS_PER_PAGE - 1) // MAX_ROWS_PER_PAGE
        if num_pages == 0: num_pages = 1

        for p in range(num_pages):
            ws = wb.copy_worksheet(master_sheet)
            ws.title = f"Page {p+1}"
            
            # --- FIX: Enforce Print Area to prevent blank pages ---
            ws.print_area = 'A1:E48'
            # ------------------------------------------------------
            
            # --- NEW: Embed QR Code on first page (Patched to B44) ---
            if p == 0 and os.path.exists(qr_img_path):
                img = XLImage(qr_img_path)
                img.width = 110
                img.height = 110
                ws.add_image(img, 'B44')
            # ---------------------------------------
            
            addons.safe_write(ws, "A4", f"Ship To : BANGOR - {outlet_name.upper()}")
            
            # --- NEW: Write Checker Display to Template ---
            addons.safe_write(ws, "A5", f"Assigned Checker: {checker_display}")
            ws["A5"].font = Font(bold=True, italic=True)
            
            addons.safe_write(ws, "D4", f"Delivery No : {delivery_no}")
            addons.safe_write(ws, "D5", f"Delivery Date : {delivery_date}")
            start_idx, end_idx = p * MAX_ROWS_PER_PAGE, (p + 1) * MAX_ROWS_PER_PAGE
            page_items = display_rows[start_idx:end_idx]
            for i, data in enumerate(page_items):
                curr_row = START_ROW + i
                ws.cell(row=curr_row, column=1, value=data["koli"]).alignment = Alignment(horizontal='center')
                ws.cell(row=curr_row, column=2, value=data["sku"])
                ws.cell(row=curr_row, column=3, value=data["qty"]).alignment = Alignment(horizontal='center')
                ws.cell(row=curr_row, column=4, value=data["uom"])
                ws.cell(row=curr_row, column=5, value=data["note"])
            
            # --- NEW: VANISH EMPTY ROWS ---
            used_rows = len(page_items)
            first_empty_row = START_ROW + used_rows
            last_template_row = START_ROW + MAX_ROWS_PER_PAGE - 1
            for r in range(first_empty_row, last_template_row + 1):
                ws.row_dimensions[r].hidden = True
            # --------------------------------
        
        wb.remove(master_sheet)
        mode_used = "Template Mode (Fixed Pagination & UOM)"
    else:
        wb = Workbook(); ws = wb.active; ws.title = "Packing List"
        
        # --- NEW: Embed QR Code (Patched to B44) ---
        if os.path.exists(qr_img_path):
            img = XLImage(qr_img_path)
            img.width = 120
            img.height = 120
            ws.add_image(img, 'B44')
        # --------------------------
        
        ws["A1"] = f"Ship To : BANGOR - {outlet_name.upper()}"
        
        # --- NEW: Write Checker Display to Default ---
        ws["A2"] = f"Checker: {checker_display}"
        ws["A2"].font = Font(bold=True, italic=True)
        
        ws["E2"], ws["F2"] = "Delivery No :", delivery_no
        ws["E3"], ws["F3"] = "Delivery Date :", delivery_date
        headers = ["No. Koli", "Description", "Qty", "Item Unit", "Notes"]
        for col, h in enumerate(headers, start=1): ws.cell(5, col, h).font = Font(bold=True)
        row = 6
        for r in display_rows:
            ws.cell(row, 1, r["koli"]); ws.cell(row, 2, r["sku"])
            ws.cell(row, 3, r["qty"]); ws.cell(row, 4, r["uom"]); ws.cell(row, 5, r["note"])
            row += 1
        mode_used = "Default Mode"

    wb.save(pl_filepath)

    # ==========================================
    #     NEW: BACKGROUND TASKS (Status, Usage, Upload)
    # ==========================================
    def _background_tasks():
        try:
            # 1. Register Packing Status (Initial state for 2-scan workflow)
            status_payload = {
                "delivery_no": delivery_no,
                "outlet": outlet_name,
                "checker": checker_display,
                "status": "PENDING", 
                "total_weight_kg": round(total_weight_kg, 2)
            }
            requests.post(f"{addons.BASE_URL}/api/packing_status", json=status_payload, headers=addons.API_HEADERS, timeout=5)
            
            # 2. Track Item Usage (for Top 25 Analytics)
            usage_payload = {
                "delivery_no": delivery_no,
                "outlet": outlet_name,
                "items": {sku: data["qty"] for sku, data in order.items()}
            }
            requests.post(f"{addons.BASE_URL}/api/track_item_usage", json=usage_payload, headers=addons.API_HEADERS, timeout=5)
            
            # 3. Upload Packing List File to Server
            if os.path.exists(pl_filepath):
                with open(pl_filepath, 'rb') as f:
                    files = {'file': (pl_filename, f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
                    data_payload = {
                        "delivery_no": delivery_no,
                        "outlet": outlet_name,
                        "checker": checker_display
                    }
                    requests.post(f"{addons.BASE_URL}/api/upload_packing_list", files=files, data=data_payload, headers=addons.API_HEADERS, timeout=10)
        except Exception:
            pass # Silent fail to not disrupt user experience
            
    threading.Thread(target=_background_tasks, daemon=True).start()
    # ==========================================

    # --- NEW: Pass Checker Display & master_db to Labels ---
    wb_labels = addons.generate_labels(outlet_name, boxes, checker_display, master_db)
    wb_labels.save(lbl_filepath)

    # --- NEW: Trigger Auto-Print Engine ---
    addons.execute_auto_print(pl_filepath, lbl_filepath)
    # --------------------------------------

    return pl_filepath, lbl_filepath, mode_used

# ==========================================
#     VAULT ENCRYPTED CRASH RECOVERY
# ==========================================

def get_recovery_file():
    if getattr(sys, 'frozen', False):
        base_dir = os.path.dirname(sys.executable)
    else:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, "logs", "recovery.dat")

def save_state():
    try:
        os.makedirs(os.path.dirname(get_recovery_file()), exist_ok=True)
        fernet = Fernet(addons.get_encryption_key())
        json_str = json.dumps(order)
        encrypted_payload = fernet.encrypt(json_str.encode())
        with open(get_recovery_file(), "wb") as f:
            f.write(encrypted_payload)
    except Exception:
        pass 

def check_recovery():
    rec_file = get_recovery_file()
    
    legacy_file = rec_file.replace(".dat", ".json")
    if os.path.exists(legacy_file):
        try: os.remove(legacy_file)
        except: pass

    if os.path.exists(rec_file):
        try:
            fernet = Fernet(addons.get_encryption_key())
            with open(rec_file, "rb") as f:
                encrypted_data = f.read()
                decrypted_data = fernet.decrypt(encrypted_data)
                saved_order = json.loads(decrypted_data.decode())
            
            if saved_order: 
                msg = "An unfinished order was detected from a previous session or crash.\n\nWould you like to restore it?"
                if messagebox.askyesno("Crash Recovery", msg):
                    order.update(saved_order)
                    update_order_display()
                else:
                    os.remove(rec_file) 
        except Exception:
            try: os.remove(rec_file)
            except: pass

def refresh_master_data():
    global master_db, DATA_SOURCE_STATE, MASTER_BOX_TOLERANCE
    global FROZEN_ITEMS, KENTANG_ITEMS, SAUCE_ITEMS, PACKAGING_ITEMS, APPAREL_ITEMS, BIG_ITEMS, BREAD_ITEMS, BUNDLE_ITEMS, DRY_ITEMS
    global ITEM_UOM, box_capacity, sku_lookup

    status_var.set("🔄 Syncing Master Data & Checkers from Cloud...")
    root.update() 

    new_db, new_state = sync_master_data()
    new_checkers = fetch_checkers()

    if new_db:
        master_db = new_db
        DATA_SOURCE_STATE = new_state

        MASTER_BOX_TOLERANCE = master_db.get("BOX_TOLERANCE", 1.859)

        FROZEN_ITEMS = set(master_db["CATEGORIES"].get("FROZEN_ITEMS", []))
        KENTANG_ITEMS = set(master_db["CATEGORIES"].get("KENTANG_ITEMS", []))
        SAUCE_ITEMS = set(master_db["CATEGORIES"].get("SAUCE_ITEMS", []))
        PACKAGING_ITEMS = set(master_db["CATEGORIES"].get("PACKAGING_ITEMS", []))
        APPAREL_ITEMS = set(master_db["CATEGORIES"].get("APPAREL_ITEMS", []))
        BIG_ITEMS = set(master_db["CATEGORIES"].get("BIG_ITEMS", []))
        BREAD_ITEMS = set(master_db["CATEGORIES"].get("BREAD_ITEMS", []))
        BUNDLE_ITEMS = set(master_db["CATEGORIES"].get("BUNDLE_ITEMS", []))

        DRY_ITEMS = (SAUCE_ITEMS | PACKAGING_ITEMS | APPAREL_ITEMS) - BUNDLE_ITEMS

        ITEM_UOM.clear()
        ITEM_UOM.update(master_db.get("ITEM_UOM", {}))

        box_capacity.clear()
        box_capacity.update(master_db.get("BOX_CAPACITY", {}))

        sku_lookup.clear()
        sku_lookup.update({sku.lower(): sku for sku in box_capacity.keys()})

        new_sku_list = list(box_capacity.keys())
        sku_dropdown['values'] = new_sku_list
        if new_sku_list:
            sku_dropdown.set(new_sku_list[0])

        # Update Checker Dropdown
        checker_combo['values'] = new_checkers
        if checker_var.get() not in new_checkers and checker_var.get() != "Select Checker":
            checker_combo.set("Select Checker")

        status_var.set(f"✅ Master Data Refreshed Live | Source: {DATA_SOURCE_STATE}")
        addons.log_action("🔄 CLOUD SYNC: Operator triggered a live Master Data refresh.")
        beep_success()
        addons.show_toast(root, f"✅ Data Refreshed: {DATA_SOURCE_STATE}", "success")
    else:
        beep_error()
        status_var.set("❌ Master Data Sync Failed.")
        addons.show_toast(root, "❌ Master Data Sync Failed.", "error")

# ==== GUI APP ====
if DND_AVAILABLE:
    root = TkinterDnD.Tk()
else:
    root = tk.Tk()

CURRENT_VERSION = "2.0.0" 
sys.app_version = CURRENT_VERSION 

addons.setup_telemetry(root)

root.title(f"Packing List Generator v{CURRENT_VERSION} Release (Created by A. Majesta P.)")

addons.apply_modern_theme(root)
root.after(100, lambda: addons.apply_button_hovers(root)) 

order, current_boxes = {}, []

status_var = tk.StringVar(value=f"🚀 System Ready | Master Data: {DATA_SOURCE_STATE}")

def add_item(event=None):
    sku, note = sku_var.get(), note_var.get().strip()
    try:
        qty = int(qty_var.get())
        if qty <= 0: raise ValueError
    except:
        beep_error()
        addons.show_toast(root, "❌ Enter valid quantity.", "error")
        return
    if not note:
        beep_error()
        addons.show_toast(root, "❌ Select a note.", "error")
        return
    
    base_qty_log = qty 
    
    if sku in ["Beef Patty Small", "Beef Patty Large"]:
        qty *= 18
    elif sku == "Thousand Island Mayonaise":
        qty *= 20
    elif sku == "Thousand Island (BBT)":
        qty *= 1
    elif sku == "Butter":
        qty *= 40

    if sku in order:
        order[sku]["qty"] += qty
        if note not in order[sku]["note"]: order[sku]["note"] = f"{order[sku]['note']}/{note}"
    else:
        order[sku] = {"qty": qty, "note": note}
    
    update_order_display()
    qty_var.set("")
    note_var.set("")
    status_var.set(f"✅ Added {qty} {sku} (Base input: {base_qty_log})")
    addons.log_action(f"➕ ADDED: {qty} units of {sku} (Note: {note})")
    beep_success()
    addons.show_toast(root, f"✅ Added {qty} {sku}", "success")

def subtract_item(event=None):
    sku = sku_var.get()
    try:
        qty = int(qty_var.get())
        if qty <= 0: raise ValueError
    except:
        beep_error()
        addons.show_toast(root, "❌ Enter valid quantity.", "error")
        return
    
    base_qty_log = qty
    
    if sku in ["Beef Patty Small", "Beef Patty Large"]:
        qty *= 18
    elif sku == "Thousand Island Mayonaise":
        qty *= 20
    elif sku == "Thousand Island (BBT)":
        qty *= 1
    elif sku == "Butter":
        qty *= 40

    if sku in order:
        order[sku]["qty"] -= qty
        if order[sku]["qty"] <= 0:
            del order[sku]
        update_order_display()
        qty_var.set("")
        status_var.set(f"➖ Removed {qty} {sku} (Base input: {base_qty_log})")
        addons.log_action(f"➖ REMOVED: {qty} units of {sku}")
        beep_success()
        addons.show_toast(root, f"➖ Removed {qty} {sku}", "success")
    else:
        beep_error()
        addons.show_toast(root, f"⚠️ {sku} not in order.", "error")

def clear_order(event=None): 
    order.clear()
    update_order_display()
    status_var.set("🗑️ Order Cleared. Ready for new input.")
    addons.log_action("🗑️ ACTION: Operator used the CLEAR ORDER button.")
    addons.show_toast(root, "🗑️ Order Cleared", "success")

def update_order_display():
    addons.refresh_order_table(order_tree, order)
    save_state() 

    if order:
        est_boxes = len(calculate_boxes(order, MASTER_BOX_TOLERANCE))
        live_est_var.set(f"📦 Live Koli Estimate: {est_boxes}")
        est_label.config(fg="#27ae60") 
    else:
        live_est_var.set("📦 Live Koli Estimate: 0")
        est_label.config(fg="#7f8c8d") 

# --- PREMIUM UPGRADE: GLOBAL LOADING CURSOR ---
def on_calculate(event=None):
    root.config(cursor="watch")
    root.update()
    try:
        global current_boxes
        if not order: 
            beep_error()
            addons.show_toast(root, "❌ Order is empty.", "error")
            return
            
        raw_boxes = calculate_boxes(order, MASTER_BOX_TOLERANCE)
        
        outlet = outlet_var.get().strip()
        if not outlet: 
            outlet = "Draft Order"
            
        def finalize_calculation(reviewed_boxes):
            global current_boxes
            current_boxes = reviewed_boxes
            status_var.set(f"🧮 Pre-Flight Approved: {len(current_boxes)} Koli ready for print.")
            addons.log_action(f"🧮 KOLI APPROVED: {len(current_boxes)} boxes ready for {outlet}.")
            beep_success()
            addons.show_toast(root, f"✅ {len(current_boxes)} Koli Approved", "success")

        addons.open_koli_reviewer(root, outlet, raw_boxes, finalize_calculation)
    finally:
        root.config(cursor="") # Always reset cursor!

def on_print(event=None):
    root.config(cursor="watch")
    root.update()
    try:
        if not current_boxes: 
            beep_error()
            addons.show_toast(root, "❌ Calculate boxes first.", "error")
            return
        outlet = outlet_var.get().strip()
        if not outlet: 
            beep_error()
            addons.show_toast(root, "❌ Enter outlet name.", "error")
            return
        
        # --- NEW: Validate or register outlet before export ---
        validated, final_outlet = validate_or_register_outlet(root, outlet)
        if not validated:
            beep_error()
            addons.show_toast(root, "❌ Export cancelled.", "error")
            return
        
        # Use the validated outlet name (might be different if user selected from suggestions)
        outlet = final_outlet
        # ------------------------------------------------------
            
        pl_fp, lbl_fp, mode = export_packing_list(outlet, current_boxes, order)
        
        addons.save_outlet_history(outlet)
        
        status_var.set(f"💾 Exported Successfully: {outlet}")
        addons.log_action(f"📤 EXPORTED: Packing List and Labels for {outlet}")
        beep_success()
        addons.show_toast(root, f"💾 Exported Successfully: {outlet}", "success")
    finally:
        root.config(cursor="") # Always reset cursor!
# --------------------------------------------

# ==========================================
#   DASHBOARD UI CONSTRUCTION (v4.4.0 UI FLIP)
# ==========================================

frame_dest, frame_entry, frame_actions, frame_feed = addons.build_dashboard_zones(root)

frame_dest.grid(row=0, column=0, sticky="nsew", padx=10, pady=(10, 5))
frame_entry.grid(row=1, column=0, sticky="nsew", padx=10, pady=5)
frame_actions.grid(row=2, column=0, sticky="nsew", padx=10, pady=5)
frame_feed.grid(row=0, column=1, rowspan=3, sticky="nsew", padx=10, pady=10)

root.grid_columnconfigure(1, weight=1)

# --- 150% DPI SCALING FIX ---
# Shift the layout flexibility so the Dropzone takes the hit instead of the input fields.
root.grid_rowconfigure(0, weight=0) # Lock Destination Details (prevents squishing)
root.grid_rowconfigure(1, weight=1) # Allow Order Entry / Dropzone to absorb the flex
root.grid_rowconfigure(2, weight=0) # Lock System Actions 

# ==========================================
#     NEW: ENHANCED DESTINATION DETAILS
# ==========================================
tk.Label(frame_dest, text="Outlet Name:", font=(addons.PREMIUM_FONT, 10, "bold"), fg="#2c3e50").grid(row=0, column=0, sticky="w", pady=5)
outlet_var = tk.StringVar()
outlet_dropdown = ttk.Combobox(frame_dest, textvariable=outlet_var)
outlet_dropdown.grid(row=0, column=1, sticky="ew", padx=10, pady=5)
frame_dest.grid_columnconfigure(1, weight=1)

# --- NEW: LIVE OUTLET PREVIEW CARD ---
preview_frame = tk.Frame(frame_dest, bg=addons.COLORS["border_light"], bd=1, relief="solid")
preview_frame.grid(row=1, column=0, columnspan=2, sticky="ew", padx=10, pady=(0, 10))

tk.Label(preview_frame, text="PREVIEW:", font=(addons.PREMIUM_FONT, 9, "bold"), 
         bg=addons.COLORS["border_light"], fg=addons.COLORS["text_primary"]).grid(row=0, column=0, sticky="w", padx=10, pady=(8, 2))

lbl_receiver = tk.Label(preview_frame, text="👤 Receiver: -", font=(addons.PREMIUM_FONT, 9), 
                        bg=addons.COLORS["border_light"], fg=addons.COLORS["text_secondary"], anchor="w")
lbl_receiver.grid(row=1, column=0, sticky="w", padx=10, pady=2)

lbl_phone = tk.Label(preview_frame, text="📞 Phone: -", font=(addons.PREMIUM_FONT, 9), 
                     bg=addons.COLORS["border_light"], fg=addons.COLORS["text_secondary"], anchor="w")
lbl_phone.grid(row=2, column=0, sticky="w", padx=10, pady=2)

lbl_address = tk.Label(preview_frame, text="📍 Address: Select an outlet to view details", 
                       font=(addons.PREMIUM_FONT, 9), bg=addons.COLORS["border_light"], 
                       fg=addons.COLORS["text_secondary"], anchor="w", wraplength=350, justify="left")
lbl_address.grid(row=3, column=0, sticky="w", padx=10, pady=(2, 8))

def update_outlet_preview(event=None):
    typed = outlet_var.get().strip().upper()
    outlet_info = master_db.get("OUTLET_INFO", {}).get(typed, {})
    if outlet_info:
        lbl_receiver.config(text=f"👤 Receiver: {outlet_info.get('name', 'N/A')}")
        lbl_phone.config(text=f"📞 Phone: {outlet_info.get('phone', 'N/A')}")
        lbl_address.config(text=f"📍 Address: {outlet_info.get('address', 'N/A')}")
    else:
        lbl_receiver.config(text="👤 Receiver: -")
        lbl_phone.config(text="📞 Phone: -")
        lbl_address.config(text="📍 Address: Select an outlet to view details")

outlet_dropdown.bind('<<ComboboxSelected>>', update_outlet_preview)
# -------------------------------------

# NEW: Checker Dropdown (Dynamic from Server)
tk.Label(frame_dest, text="Assigned Checker:", font=(addons.PREMIUM_FONT, 10, "bold"), fg="#2c3e50").grid(row=2, column=0, sticky="w", pady=5)
checker_var = tk.StringVar()
checker_combo = ttk.Combobox(frame_dest, textvariable=checker_var, state="readonly", font=(addons.PREMIUM_FONT, 11))

# Fetch dynamic checkers from server, fallback to hardcoded if offline
dynamic_checkers = fetch_checkers()
checker_combo['values'] = dynamic_checkers
checker_combo.set("Select Checker")
checker_combo.grid(row=2, column=1, sticky="ew", padx=10, pady=5)

# NEW: Cluster Textbox
tk.Label(frame_dest, text="Cluster / Route:", font=(addons.PREMIUM_FONT, 10, "bold"), fg="#2c3e50").grid(row=3, column=0, sticky="w", pady=5)
cluster_var = tk.StringVar()
cluster_entry = tk.Entry(frame_dest, textvariable=cluster_var, font=(addons.PREMIUM_FONT, 11))
cluster_entry.grid(row=3, column=1, sticky="ew", padx=10, pady=5)

# Initialize dropdown with master data outlets + history for fuzzy search
outlet_history = addons.load_outlet_history()

def update_outlet_suggestions(event):
    # SAFE CHECK: plain tk.Event() objects (used for initial load) have no 'keysym'
    if getattr(event, "keysym", None) in ('Up', 'Down', 'Left', 'Right', 'Return'):
        return
    typed = outlet_var.get().upper()
    
    # Combine master data and history dynamically
    current_master_outlets = list(master_db.get("OUTLET_INFO", {}).keys())
    current_history = addons.load_outlet_history()
    combined_outlets = list(set(current_master_outlets + current_history))
    
    if typed == "":
        outlet_dropdown['values'] = sorted(combined_outlets)
    else:
        filtered = [item for item in combined_outlets if typed in item.upper()]
        outlet_dropdown['values'] = sorted(filtered)
    update_outlet_preview(event) # Update preview as they type/select
    
outlet_dropdown.bind('<KeyRelease>', update_outlet_suggestions)
# Trigger initial load
update_outlet_suggestions(tk.Event())
# ==========================================

# --- NEW: LOGGED-IN OPERATOR BADGE ---
operator_alias = addons.HWID_ALIASES.get(addons.get_hwid(), "Operator")
logged_in_label = tk.Label(frame_feed, text=f"👤 Logged In As: {operator_alias}", 
                           font=(addons.PREMIUM_FONT, 10, "bold"), fg=addons.COLORS["accent_purple"],
                           bg=addons.COLORS["bg_primary"])
logged_in_label.pack(anchor="e", pady=(0, 5))
# -------------------------------------

live_est_var = tk.StringVar()
live_est_var.set("📦 Live Koli Estimate: 0")
est_label = tk.Label(frame_feed, textvariable=live_est_var, font=(addons.PREMIUM_FONT, 11, "bold"), 
                     fg=addons.COLORS["text_secondary"], bg=addons.COLORS["bg_primary"])
est_label.pack(anchor="e", pady=(0, 10))

table_container = tk.Frame(frame_feed)
table_container.pack(fill="both", expand=True)
order_tree, order_scroll = addons.build_order_table(table_container)
order_tree.pack(side="left", fill="both", expand=True)
order_scroll.pack(side="right", fill="y")
addons.enable_quick_edit(root, order_tree, order, update_order_display)

addons.enable_context_menu(root, order_tree, order, update_order_display)


if DND_AVAILABLE:
    tabs = ttk.Notebook(frame_entry)
    tabs.pack(fill="both", expand=True, pady=2, padx=2)

    tab_scan = tk.Frame(tabs, bg=addons.COLORS["bg_primary"])
    tab_manual = tk.Frame(tabs, bg=addons.COLORS["bg_primary"])

    tabs.add(tab_scan, text="🚀 AUTO-DROPZONE")
    tabs.add(tab_manual, text="🛠️ MANUAL OVERRIDE")
    
    drop_label = tk.Label(tab_scan, text="📄 DRAG & DROP\nPDF HERE", 
                          font=(addons.PREMIUM_FONT, 14, "bold"), bg=addons.COLORS["border_light"], 
                          fg=addons.COLORS["text_primary"], relief="sunken", bd=4)
    drop_label.pack(fill="both", expand=True, padx=15, pady=15)

    # --- PREMIUM UPGRADE: INTERACTIVE DRAG & DROP FEEDBACK ---
    def on_drag_enter(event):
        drop_label.config(bg=addons.COLORS["accent_green"], fg="white", text="📥 DROP FILE HERE")

    def on_drag_leave(event):
        drop_label.config(bg=addons.COLORS["border_light"], fg=addons.COLORS["text_primary"], text="📄 DRAG & DROP\nPDF HERE")

    drop_label.dnd_bind("<<DragEnter>>", on_drag_enter)
    drop_label.dnd_bind("<<DragLeave>>", on_drag_leave)
    # --------------------------------------------------------

    def handle_drop(event):
        files = root.tk.splitlist(event.data)
        success_count = 0
        
        pending_totals = {}
        processed_results_list = []
        
        for file_path in files:
            clean_path = file_path.strip('{}')
            if not os.path.isfile(clean_path): continue
            
            results = None
            if clean_path.lower().endswith(('.xlsx', '.xls')):
                results = addons.smart_scan_file(clean_path, sku_lookup)
            elif clean_path.lower().endswith('.pdf'):
                results = addons.smart_scan_pdf(clean_path, sku_lookup)
                
            if results is not None:
                success_count += 1
                processed_results_list.append(results)
                
                # Pre-calculate totals for validation
                for sku, data in results.items():
                    adjusted_qty = data["qty"]
                    if sku in ["Beef Patty Small", "Beef Patty Large"]: adjusted_qty *= 18
                    elif sku == "Thousand Island Mayonaise": adjusted_qty *= 20
                    elif sku == "Thousand Island (BBT)": adjusted_qty *= 1
                    elif sku == "Butter": adjusted_qty *= 40

                    if sku in pending_totals:
                        pending_totals[sku] += adjusted_qty
                    else:
                        pending_totals[sku] = adjusted_qty

        if success_count > 0:
            # --- NEW: Cloud Stock Interception ---
            current_stock = addons.sync_current_stock()
            is_safe = addons.validate_stock_levels(root, pending_totals, current_stock)

            if is_safe:
                for results in processed_results_list:
                    for sku, data in results.items():
                        adjusted_qty = data["qty"]
                        if sku in ["Beef Patty Small", "Beef Patty Large"]: adjusted_qty *= 18
                        elif sku == "Thousand Island Mayonaise": adjusted_qty *= 20
                        elif sku == "Thousand Island (BBT)": adjusted_qty *= 1
                        elif sku == "Butter": adjusted_qty *= 40

                        if sku in order:
                            order[sku]["qty"] += adjusted_qty
                            if data["note"] not in order[sku]["note"]:
                                order[sku]["note"] += f"/{data['note']}"
                        else:
                            order[sku] = {"qty": adjusted_qty, "note": data["note"]}

                update_order_display()
                status_var.set(f"🎯 Drop-Scanned successfully: {success_count} file(s).")
                addons.log_action(f"🎯 SCANNED: Drag & Drop success for {success_count} file(s).")
                beep_success()
                addons.show_toast(root, f"🎯 Scanned {success_count} file(s) successfully.", "success")
            else:
                status_var.set("⚠️ Import Cancelled due to stock shortages.")
                addons.log_action("⚠️ CANCELLED: Drop scan aborted by operator due to stock shortages.")
                beep_error()
                addons.show_toast(root, "⚠️ Import Cancelled", "error")
        else:
            status_var.set("❌ Invalid file dropped or scan failed.")
            addons.log_action("❌ ERROR: Drag & Drop scan failed or invalid file.")
            beep_error()
            addons.show_toast(root, "❌ Drop scan failed.", "error")

    drop_label.drop_target_register(DND_FILES)
    drop_label.dnd_bind('<<Drop>>', handle_drop)
    
    manual_parent = tab_manual
else:
    manual_parent = frame_entry
    tk.Label(frame_entry, text="⚠️ Install 'tkinterdnd2' via pip to enable the Drag-and-Drop zone.", fg="red").grid(row=10, column=0, columnspan=2)

filter_frame = tk.Frame(manual_parent)
filter_frame.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 8))

def apply_filter(cat_name, item_set=None):
    if item_set:
        filtered = [sku for sku in box_capacity.keys() if sku in item_set]
        sku_dropdown['values'] = filtered
        if filtered: sku_dropdown.set(filtered[0])
        else: sku_dropdown.set("")
    else:
        sku_dropdown['values'] = list(box_capacity.keys())
        sku_dropdown.set("Beef Patty Small")
    status_var.set(f"🔍 Filter Applied: {cat_name}")

tk.Button(filter_frame, text="All", font=(addons.PREMIUM_FONT, 8), command=lambda: apply_filter("All")).pack(side="left", padx=2)
tk.Button(filter_frame, text="❄️ Frozen", font=(addons.PREMIUM_FONT, 8), bg="#3498db", fg="white", command=lambda: apply_filter("Frozen", FROZEN_ITEMS | KENTANG_ITEMS)).pack(side="left", padx=2)
tk.Button(filter_frame, text="🥫 Sauce & Bread", font=(addons.PREMIUM_FONT, 8), bg="#e67e22", fg="white", command=lambda: apply_filter("Sauce/Bread", SAUCE_ITEMS | BREAD_ITEMS)).pack(side="left", padx=2)
tk.Button(filter_frame, text="📦 Pack", font=(addons.PREMIUM_FONT, 8), bg="#9b59b6", fg="white", command=lambda: apply_filter("Packaging", PACKAGING_ITEMS | BUNDLE_ITEMS)).pack(side="left", padx=2)
tk.Button(filter_frame, text="👕 Merch", font=(addons.PREMIUM_FONT, 8), bg="#34495e", fg="white", command=lambda: apply_filter("Merch", APPAREL_ITEMS)).pack(side="left", padx=2)

tk.Label(manual_parent, text="Select SKU:").grid(row=1, column=0, sticky="w", pady=2)
sku_var = tk.StringVar()
sku_dropdown = ttk.Combobox(manual_parent, textvariable=sku_var, values=list(box_capacity.keys()))
sku_dropdown.grid(row=1, column=1, sticky="ew", padx=10, pady=2)
sku_dropdown.set("Beef Patty Small")

tk.Label(manual_parent, text="Quantity:").grid(row=2, column=0, sticky="w", pady=2)
qty_var = tk.StringVar()
tk.Entry(manual_parent, textvariable=qty_var).grid(row=2, column=1, sticky="ew", padx=10, pady=2)

tk.Label(manual_parent, text="Note:").grid(row=3, column=0, sticky="w", pady=2)
note_var = tk.StringVar()
note_dropdown = ttk.Combobox(manual_parent, textvariable=note_var, values=["BGB", "BBB", "✓"], state="readonly")
note_dropdown.grid(row=3, column=1, sticky="ew", padx=10, pady=2)

manual_parent.grid_columnconfigure(1, weight=1)

btn_frame = tk.Frame(manual_parent)
btn_frame.grid(row=4, column=0, columnspan=2, pady=10, sticky="ew") # Shifted to Row 4
btn_frame.grid_columnconfigure(0, weight=1)
btn_frame.grid_columnconfigure(1, weight=1)
tk.Button(btn_frame, text="+ Add Item", command=add_item, bg="#2ecc71", fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=0, column=0, sticky="ew", padx=(0, 5))
tk.Button(btn_frame, text="- Subtract Item", command=subtract_item, bg="#e74c3c", fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=0, column=1, sticky="ew", padx=(5, 0))

# ==========================================
#     NEW: SIMPLIFIED SYSTEM ACTIONS
# ==========================================# Row 0: Primary Actions
tk.Button(frame_actions, text="🧮 Calculate Routing", command=on_calculate, 
          bg=addons.COLORS["accent_blue"], fg="white", font=(addons.PREMIUM_FONT, 10, "bold")).grid(row=0, column=0, sticky="ew", padx=5, pady=5)
tk.Button(frame_actions, text="💾 Export & Save Data", command=on_print, 
          bg=addons.COLORS["accent_green"], fg="white", font=(addons.PREMIUM_FONT, 10, "bold")).grid(row=0, column=1, sticky="ew", padx=5, pady=5)

# Row 1: Tools Grid
tk.Button(frame_actions, text="📂 Scanner", command=lambda: addons.open_auto_mode(root, order, sku_lookup, update_order_display), 
          bg=addons.COLORS["accent_purple"], fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=1, column=0, sticky="ew", padx=5, pady=2)
tk.Button(frame_actions, text="🔄 Live Sync Data", command=refresh_master_data, 
          bg="#16a085", fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=1, column=1, sticky="ew", padx=5, pady=2)

# Row 2: Tools Grid
tk.Button(frame_actions, text="🖨️ Printer Settings", command=lambda: addons.open_printer_settings(root), 
          bg="#95a5a6", fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=2, column=0, sticky="ew", padx=5, pady=2)
tk.Button(frame_actions, text="📊 Shift Report", command=lambda: addons.view_shift_report(root), 
          bg=addons.COLORS["accent_orange"], fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=2, column=1, sticky="ew", padx=5, pady=2)

# Row 3: Admin & Extras
tk.Button(frame_actions, text="💳 Wallet", command=lambda: addons.open_wallet_ui(root, addons.get_hwid()), 
          bg=addons.COLORS["accent_purple"], fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=3, column=0, sticky="ew", padx=5, pady=2)
tk.Button(frame_actions, text="📡 Live Board", command=lambda: addons.open_packing_board(root), 
          bg=addons.COLORS["accent_blue"], fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=3, column=1, sticky="ew", padx=5, pady=2)

# Row 4: Updates
tk.Button(frame_actions, text="🔄 Check for Updates", command=lambda: addons.check_for_updates(root, CURRENT_VERSION, silent=False), 
          bg=addons.COLORS["accent_orange"], fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=4, column=0, columnspan=2, sticky="ew", padx=5, pady=5)

# Row 5: Danger Zone
tk.Button(frame_actions, text="🗑️ CLEAR CURRENT ORDER", command=clear_order, 
          bg=addons.COLORS["accent_red"], fg="white", font=(addons.PREMIUM_FONT, 9, "bold")).grid(row=5, column=0, columnspan=2, sticky="ew", padx=5, pady=(5, 10))

frame_actions.grid_columnconfigure(0, weight=1)
frame_actions.grid_columnconfigure(1, weight=1)
# ==========================================

status_bar = tk.Label(root, textvariable=status_var, bd=1, relief=tk.SUNKEN, anchor="w", 
                      font=(addons.PREMIUM_FONT, 9, "italic"), fg=addons.COLORS["text_primary"], 
                      padx=10, bg=addons.COLORS["border_light"])
status_bar.grid(row=4, column=0, columnspan=2, sticky="ew")

root.bind('<Return>', add_item)             
root.bind('<Control-Return>', on_calculate) 
root.bind('<Control-s>', on_print)          
root.bind('<Control-x>', clear_order)       
root.bind('<Delete>', clear_order)          

addons.apply_scaling(root)
root.after(100, lambda: addons.validate_license(root))

hwid = addons.get_hwid() 
addons.start_live_telemetry(hwid, status_var)
addons.start_shadow_logger(hwid)
addons.log_action("⚡ SYSTEM BOOT: Terminal Online & Overwatch Connected.")

def safe_start_checks():
    if root.state() == 'normal':
        check_recovery()
        
        # --- NEW: First-Time Printer Setup Check ---
        if not os.path.exists(os.path.join("logs", "printer_settings.json")):
            root.after(1500, lambda: addons.open_printer_settings(root))
        # -------------------------------------------

        root.after(2000, lambda: addons.check_for_updates(root, CURRENT_VERSION, silent=True))
    else:
        root.after(500, safe_start_checks)

root.after(500, safe_start_checks) 

root.mainloop()
