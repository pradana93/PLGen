import tkinter as tk
from tkinter import messagebox, ttk, filedialog, simpledialog
import datetime
import sys
import os
import json
import requests
import time
import re
import subprocess
import uuid
import hashlib
from email.utils import parsedate_to_datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Side, Border, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
import pandas as pd
import pdfplumber
import signal
import threading
import traceback
import copy
from cryptography.fernet import Fernet
import base64
import random # ✨ Required for Ghost Math ✨
# --- NEW IMPORTS FOR EMAIL ENGINE ---
import smtplib
from email.message import EmailMessage
import mimetypes
# ------------------------------------
# --- NEW IMPORTS FOR WALLET & QRIS ---
import qrcode
from PIL import Image, ImageTk
# -------------------------------------

# --- NEW: PREMIUM UI CONFIGURATION ---
PREMIUM_FONT = "Segoe UI"

# 🎨 COLOR PALETTE - MODERN THEME
COLORS = {
    "bg_primary": "#f4f6f9",        # Light gray-blue background
    "bg_secondary": "#ffffff",      # Pure white cards
    "bg_dark": "#2c3e50",           # Dark blue-gray for headers
    "accent_blue": "#3498db",       # Primary action color
    "accent_green": "#27ae60",      # Success actions
    "accent_red": "#e74c3c",        # Error/danger
    "accent_orange": "#f39c12",     # Warnings
    "accent_purple": "#9b59b6",     # Premium features
    "text_primary": "#2c3e50",      # Main text
    "text_secondary": "#7f8c8d",    # Muted text
    "border_light": "#ecf0f1",      # Subtle borders
    "shadow": "rgba(0,0,0,0.1)",    # Shadow color
    
    # Category Colors
    "frozen": "#3498db",            # Blue for frozen items
    "kentang": "#f39c12",           # Orange for potato items
    "sauce": "#e74c3c",             # Red for sauces
    "packaging": "#27ae60",         # Green for packaging
    "apparel": "#9b59b6",           # Purple for apparel
    "big": "#e67e22",               # Dark orange for big items
    "bread": "#f1c40f",             # Yellow for bread
    "bundle": "#1abc9c",            # Teal for bundles
}

# DARK MODE COLORS
DARK_COLORS = {
    "bg_primary": "#1a1a2e",        # Dark navy background
    "bg_secondary": "#16213e",      # Darker cards
    "bg_dark": "#0f3460",           # Header background
    "accent_blue": "#4fc3f7",       # Lighter blue for dark mode
    "accent_green": "#81c784",      # Softer green
    "accent_red": "#e57373",        # Softer red
    "accent_orange": "#ffb74d",     # Softer orange
    "accent_purple": "#ba68c8",     # Softer purple
    "text_primary": "#eceff1",      # Light text
    "text_secondary": "#b0bec5",    # Muted light text
    "border_light": "#2c3e50",      # Darker borders
    "shadow": "rgba(0,0,0,0.3)",    
}

# Global theme state
CURRENT_THEME = "light"  # "light" or "dark"
# -------------------------------------

# --- NEW: Company Switch ---
ACTIVE_COMPANY_CODE = "BBB"
_LAST_SCANNED_FILE = None  # <-- NEW: Memory stash for Drag & Drop files
# ---------------------------

# ==========================================
#     GLOBAL SECURITY CONFIGURATION
# ==========================================
# API Token to prevent unauthorized database spam
API_HEADERS = {"Authorization": "Bearer JESTA-SECURE-99X"}
BASE_URL = "https://jestu93.pythonanywhere.com"

# ✨ 5. HWID Aliases for Splash Screen ✨
HWID_ALIASES = {
    "8DB7CE3731E42814": "Majesta (Lead Developer)",
    "A958AAA787BF6FF9": "Zahra Logistic VT",
    "30EA5F1E9BD8FD68": "Nur Logistic VT"
}

# JAC (Jesta Anti-Cheat) Master Flag
PIRATE_MODE = False

# --- NEW: OFFLINE MODE SECURITY ---
OFFLINE_MODE = False
MAJESTA_SECRET_SALT = "JESTA_OFFLINE_VAULT_2026"
_consecutive_failures = 0

def generate_offline_pin(hwid, date_str):
    """Generates a 6-digit daily PIN based on HWID and Date."""
    raw = f"{hwid}|{date_str}|{MAJESTA_SECRET_SALT}"
    hash_val = hashlib.sha256(raw.encode()).hexdigest()
    digits = ''.join(filter(str.isdigit, hash_val))
    return digits[:6].ljust(6, '0')

def check_offline_pin(entered_pin, hwid):
    """Validates PIN against Today and Yesterday to prevent midnight lockouts."""
    today_str = datetime.datetime.now().strftime("%Y%m%d")
    yesterday_str = (datetime.datetime.now() - datetime.timedelta(days=1)).strftime("%Y%m%d")
    
    if entered_pin == generate_offline_pin(hwid, today_str): return True
    if entered_pin == generate_offline_pin(hwid, yesterday_str): return True
    return False

# --- NEW: ENCRYPTED LICENSE VAULT HELPERS ---
def get_vault_path():
    if getattr(sys, 'frozen', False): base_dir = os.path.dirname(sys.executable)
    else: base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, "logs", "license_vault.dat")

def load_vault():
    path = get_vault_path()
    if not os.path.exists(path): return None
    try:
        fernet = Fernet(get_encryption_key())
        with open(path, "rb") as f:
            return json.loads(fernet.decrypt(f.read()).decode())
    except:
        return None

def save_vault(data):
    path = get_vault_path()
    try:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        fernet = Fernet(get_encryption_key())
        with open(path, "wb") as f:
            f.write(fernet.encrypt(json.dumps(data).encode()))
    except: pass
# ------------------------------------------

# ✨ 3. The Checksum Trap ✨
def run_jac_scan():
    """Hashes the executable on boot. If modified, silently arms the Ghost Math & Tattletale."""
    global PIRATE_MODE
    try:
        target_file = sys.executable if getattr(sys, 'frozen', False) else "core.py"
        with open(target_file, "rb") as f:
            file_hash = hashlib.sha256(f.read()).hexdigest()
            
        # ⚠️ DEPLOYMENT STEP: Once compiled with Nuitka, hash your .exe and paste it here!
        MASTER_HASH = "REPLACE_ME_WITH_REAL_EXE_HASH" 
        
        if MASTER_HASH != "REPLACE_ME_WITH_REAL_EXE_HASH" and file_hash != MASTER_HASH:
            PIRATE_MODE = True
    except Exception:
        pass # Fail silently so actual users don't crash if an AV locks the file read

# ==========================================
#     NEW: PREMIUM MODERN INPUT FIELD
# ==========================================
class ModernEntry(tk.Frame):
    """Custom modern input field with focus-in/out border color change."""
    def __init__(self, parent, textvariable=None, font=(PREMIUM_FONT, 11), width=20, **kwargs):
        super().__init__(parent, bg="#bdc3c7", bd=0, highlightthickness=1, highlightbackground="#bdc3c7", **kwargs)
        self.entry = tk.Entry(self, textvariable=textvariable, font=font, bd=0, highlightthickness=0, width=width, bg="white")
        self.entry.pack(padx=1, pady=1, fill="both", expand=True)
        
        self.entry.bind("<FocusIn>", lambda e: self.config(highlightbackground="#2980b9", highlightthickness=2))
        self.entry.bind("<FocusOut>", lambda e: self.config(highlightbackground="#bdc3c7", highlightthickness=1))
        
    def get(self): return self.entry.get()
    def insert(self, index, text): self.entry.insert(index, text)
    def delete(self, first, last=None): self.entry.delete(first, last)
    def bind(self, *args, **kwargs): self.entry.bind(*args, **kwargs)

# ==========================================
#         LICENSE & SECURITY (CLOUD v3.0)
# ==========================================

def get_hwid():
    try:
        cmd = subprocess.Popen("wmic csproduct get uuid", shell=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        output, _ = cmd.communicate()
        sys_uuid = output.decode('utf-8').split('\n')[1].strip()
        
        if not sys_uuid or sys_uuid == "FFFFFFFF-FFFF-FFFF-FFFF-FFFFFFFFFFFF" or "ERROR" in sys_uuid.upper():
            raise ValueError("Invalid WMI UUID")
            
        return hashlib.md5(sys_uuid.encode()).hexdigest()[:16].upper()
    
    except Exception:
        mac = uuid.getnode()
        return hashlib.md5(str(mac).encode()).hexdigest()[:16].upper()

def get_true_network_time():
    try:
        response = requests.head("https://www.google.com", timeout=3)
        date_str = response.headers['Date']
        return parsedate_to_datetime(date_str).replace(tzinfo=None)
    except:
        return None

# --- NEW: OFFLINE STARTUP GATE (STRICT RULES) ---
def offline_startup_gate(root):
    current_hwid = get_hwid()
    vault = load_vault()
    today_str = datetime.datetime.now().strftime("%Y%m%d")
    
    # Rule 1: No virgin machine
    if not vault:
        messagebox.showerror("Offline Boot Denied", "Internet connection required for first-time activation.\nPlease connect to the internet and restart.", parent=root)
        sys.exit()
        
    # Rule 4: Clock tamper detection
    last_online = vault.get("last_online_date", "19700101")
    if today_str < last_online:
        messagebox.showerror("Security Alert", "System clock tampering detected.\nOffline mode disabled.", parent=root)
        sys.exit()
        
    # Rule 2: Cached license gate
    if vault.get("last_license_status") != "ACTIVE":
        messagebox.showerror("Offline Boot Denied", "Your last known license status is not ACTIVE.\nPlease connect to the internet to renew.", parent=root)
        sys.exit()
        
    # Rule 2 & Expiry Proximity Lock
    try:
        expiry_dt = datetime.datetime.strptime(vault.get("license_expiry", "19700101"), "%Y%m%d")
        today_dt = datetime.datetime.strptime(today_str, "%Y%m%d")
        days_left = (expiry_dt - today_dt).days
        if days_left < 0:
            messagebox.showerror("Offline Boot Denied", "Your cached license has expired.\nPlease connect to the internet to renew.", parent=root)
            sys.exit()
        if days_left <= 7:
            messagebox.showerror("Offline Boot Denied", "Your license expires in less than 7 days.\nOffline mode is locked. Please renew online.", parent=root)
            sys.exit()
    except:
        pass
        
    # Rule 3: Offline streak limit
    streak = vault.get("offline_streak", 0)
    if streak >= 3:
        messagebox.showerror("Offline Boot Denied", "Maximum offline streak (3 days) reached.\nYou must connect to the internet to verify your license.", parent=root)
        sys.exit()
        
    # Rule 5: PIN attempts
    pin_tracker = vault.get("pin_attempts_today", {})
    if pin_tracker.get("date") != today_str:
        pin_tracker = {"date": today_str, "count": 0}
    if pin_tracker["count"] >= 3:
        messagebox.showerror("Offline Boot Denied", "Too many failed PIN attempts today.\nOffline mode locked until tomorrow.", parent=root)
        sys.exit()

    # Show UI
    gate_win = tk.Toplevel(root)
    gate_win.title("📡 Offline Startup Request")
    gate_win.geometry("450x420")
    gate_win.attributes("-topmost", True)
    gate_win.overrideredirect(True)
    gate_win.configure(bg="#2c3e50")
    
    gate_win.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() // 2) - 225
    y = root.winfo_y() + (root.winfo_height() // 2) - 210
    gate_win.geometry(f"+{x}+{y}")
    
    tk.Label(gate_win, text="📡", font=("Segoe UI Emoji", 40), bg="#2c3e50", fg="white").pack(pady=(20, 5))
    tk.Label(gate_win, text="OFFLINE STARTUP REQUEST", font=(PREMIUM_FONT, 14, "bold"), bg="#2c3e50", fg="#f1c40f").pack()
    tk.Label(gate_win, text=f"Terminal ID: {current_hwid}\nOffline Day {streak + 1} of 3", font=(PREMIUM_FONT, 10), bg="#2c3e50", fg="#ecf0f1", justify="center").pack(pady=5)
    
    tk.Label(gate_win, text="Enter 6-digit PIN from Lead Dev to boot offline:", font=(PREMIUM_FONT, 10), bg="#2c3e50", fg="#bdc3c7").pack(pady=(10, 0))
    
    pin_var = tk.StringVar()
    pin_entry = tk.Entry(gate_win, textvariable=pin_var, font=(PREMIUM_FONT, 16, "bold"), justify="center", width=10, bg="#34495e", fg="white", insertbackground="white")
    pin_entry.pack(pady=10)
    
    error_lbl = tk.Label(gate_win, text="", font=(PREMIUM_FONT, 10, "bold"), bg="#2c3e50", fg="#e74c3c")
    error_lbl.pack()
    
    def attempt_boot():
        nonlocal pin_tracker
        entered = pin_var.get().strip()
        if check_offline_pin(entered, current_hwid):
            global OFFLINE_MODE
            OFFLINE_MODE = True
            vault["offline_streak"] = streak + 1
            vault["last_offline_date"] = today_str
            save_vault(vault)
            
            # Rule 6: Audit trail (queued locally since we are offline)
            try:
                queue_file = os.path.join("logs", "pending_sync.json")
                queue = []
                if os.path.exists(queue_file):
                    with open(queue_file, "r") as f: queue = json.load(f)
                queue.append({"endpoint": "/api/audit_logs", "payload": {
                    "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "user": HWID_ALIASES.get(current_hwid, current_hwid),
                    "role": "Terminal",
                    "action_type": "OFFLINE_BOOT",
                    "details": f"Booted offline. Streak: {streak + 1}/3"
                }})
                with open(queue_file, "w") as f: json.dump(queue, f)
            except: pass
            
            gate_win.destroy()
            root.deiconify()
            root.update()
            # Start heartbeat (it will run in offline mode and queue/try to reconnect)
            start_heartbeat(root)
            show_toast(root, f"🔓 OFFLINE MODE (Day {streak + 1}/3) - Local Cache Active", "success")
            log_action(f"🔓 OFFLINE BOOT: Authorized via PIN. Day {streak + 1}/3.")
        else:
            pin_tracker["count"] += 1
            vault["pin_attempts_today"] = pin_tracker
            save_vault(vault)
            remaining = 3 - pin_tracker["count"]
            if remaining <= 0:
                error_lbl.config(text="❌ Too many failed attempts. Locked for today.")
                pin_entry.config(state="disabled")
            else:
                error_lbl.config(text=f"❌ Invalid PIN. {remaining} attempt(s) left today.")
            pin_var.set("")
            
    tk.Button(gate_win, text="🔓 BOOT OFFLINE", command=attempt_boot, bg="#f39c12", fg="white", font=(PREMIUM_FONT, 11, "bold"), width=15).pack(pady=10)
    
    root.wait_window(gate_win)
    if not OFFLINE_MODE:
        sys.exit() # If they closed the window or failed 3 times
# ------------------------------------------------

def validate_license(root):
    CLOUD_URL = f"{BASE_URL}/check"
    current_hwid = get_hwid()
    
    # Initialize Anti-Cheat Bootloader
    run_jac_scan()

    root.withdraw() 

    splash = tk.Toplevel(root)
    splash.title("Authenticating")
    splash.minsize(550, 400) 
    splash.geometry("550x400")
    splash.attributes("-topmost", True)

    container = tk.Frame(splash, padx=30, pady=30)
    container.pack(expand=True, fill="both")
    
    tk.Label(container, text="License Check", font=(PREMIUM_FONT, 18, "bold")).pack(pady=(10, 20))
    
    status_label = tk.Label(container, text=f"Verifying ID: {current_hwid}\nConnecting to secure server...", 
                            font=(PREMIUM_FONT, 12), wraplength=450, justify="center")
    status_label.pack(pady=10)

    progress = ttk.Progressbar(container, orient="horizontal", length=400, mode="determinate")
    progress.pack(pady=20)
    
    countdown_label = tk.Label(container, text="", font=(PREMIUM_FONT, 13, "bold"))
    countdown_label.pack(pady=10)
    
    #  4. Transparency Disclaimer (HR & Legal Shield) 
    disclaimer = "🔒 Protected by Jesta Cloud. Terminal telemetry and document scans are monitored for Quality Assurance."
    tk.Label(container, text=disclaimer, font=(PREMIUM_FONT, 8), fg="#7f8c8d", wraplength=450).pack(side="bottom", pady=(20, 0))
    
    splash.update()

    for i in range(1, 41):
        progress['value'] = i
        splash.update()
        time.sleep(0.01)

    try:
        response = requests.get(CLOUD_URL, params={"id": current_hwid}, headers=API_HEADERS, timeout=8)
        true_time = get_true_network_time()
        
        for i in range(41, 101):
            progress['value'] = i
            splash.update()
            time.sleep(0.005)

        if not true_time:
            messagebox.showerror("Time Sync Error", "Could not verify true network time.\nPlease check your internet connection.", parent=splash)
            sys.exit()

        if response.status_code == 200:
            data = response.json()
            if data.get("status") == "ACTIVE":
                expiry_str = data.get("expiry")
                try:
                    expiry_date = datetime.datetime.strptime(expiry_str, "%Y%m%d")
                    days_left = (expiry_date - true_time).days + 1
                    
                    if days_left < 0:
                        messagebox.showerror("Expired", "Subscription expired.\nContact Jesta.", parent=splash)
                        sys.exit()

                    color = "red" if days_left <= 5 else "green"
                    
                    # ✨ 5. Display the Welcome Message ✨
                    alias = HWID_ALIASES.get(current_hwid, "Operator")
                    status_label.config(text=f"ACCESS GRANTED\nWelcome back, {alias}!", fg="green", font=(PREMIUM_FONT, 14, "bold"))
                    countdown_label.config(text=f"Subscription: {days_left} Days Remaining", fg=color)
                    
                    splash.update()
                    time.sleep(2.0)
                    
                except Exception:
                    alias = HWID_ALIASES.get(current_hwid, "Operator")
                    status_label.config(text=f"ACCESS GRANTED (Lifetime)\nWelcome back, {alias}!", fg="green", font=(PREMIUM_FONT, 14, "bold"))
                    expiry_str = "20990101" # Fallback for lifetime
                    splash.update()
                    time.sleep(2.0)

                # --- NEW: SAVE VAULT ON SUCCESS ---
                today_str = true_time.strftime("%Y%m%d") if true_time else datetime.datetime.now().strftime("%Y%m%d")
                vault_data = {
                    "last_online_date": today_str,
                    "last_license_status": "ACTIVE",
                    "license_expiry": expiry_str,
                    "offline_streak": 0, # Reset streak
                    "pin_attempts_today": {"date": today_str, "count": 0}
                }
                # Preserve pin attempts if it's the same day
                old_vault = load_vault()
                if old_vault and old_vault.get("pin_attempts_today", {}).get("date") == today_str:
                    vault_data["pin_attempts_today"] = old_vault["pin_attempts_today"]
                save_vault(vault_data)
                # ----------------------------------

                splash.destroy()
                root.deiconify() 
                root.update()
                start_heartbeat(root)
                return True
            else:
                splash.destroy()
                # --- PATCHED: REDIRECT TO WALLET UI ---
                # We pass force_exit=True because they are blocked at the login screen
                open_wallet_ui(root, current_hwid, force_exit=True)
                sys.exit()
                # --------------------------------------
        else:
            messagebox.showerror("Server Error", "Cloud Gatekeeper returned an error.", parent=splash)
            sys.exit()

    except requests.exceptions.RequestException:
        # --- NEW: TRIGGER OFFLINE GATE INSTEAD OF EXIT ---
        if splash and splash.winfo_exists():
            splash.destroy()
        offline_startup_gate(root)
        # --------------------------------------------------

def show_registration_info(root, hwid):
    # Legacy fallback, now redirects to wallet
    open_wallet_ui(root, hwid, force_exit=True)
    sys.exit()

# --- NEW: PROFESSIONAL CONNECTION SHIELD (WITH OFFLINE PIN) ---
_OFFLINE_LOCK_ACTIVE = False
_lock_window = None

def show_offline_lock(root):
    """Deploys a professional, native gray-out shield when the internet drops."""
    global _OFFLINE_LOCK_ACTIVE, _lock_window
    if _OFFLINE_LOCK_ACTIVE: return
    _OFFLINE_LOCK_ACTIVE = True

    # 1. Natively disable the main app (Greys it out and blocks all clicks instantly)
    try:
        root.attributes('-disabled', True)
    except:
        pass

    # 2. Create a clean, modern modal dialog
    _lock_window = tk.Toplevel(root)
    _lock_window.title("Connection Lost")
    _lock_window.geometry("450x420") # Increased height for PIN entry
    _lock_window.attributes("-topmost", True)
    _lock_window.overrideredirect(True) # Clean borderless look
    _lock_window.configure(bg="#2c3e50")
    
    # Center it over the app
    _lock_window.update_idletasks()
    x = root.winfo_x() + (root.winfo_width() // 2) - 225
    y = root.winfo_y() + (root.winfo_height() // 2) - 210
    _lock_window.geometry(f"+{x}+{y}")

    # UI Elements
    tk.Label(_lock_window, text="📡", font=("Segoe UI Emoji", 40), bg="#2c3e50", fg="white").pack(pady=(20, 10))
    
    tk.Label(_lock_window, text="CONNECTION LOST", font=(PREMIUM_FONT, 16, "bold"), bg="#2c3e50", fg="#e74c3c").pack()
    
    current_hwid = get_hwid()
    tk.Label(_lock_window, text=f"Terminal ID: {current_hwid}", font=(PREMIUM_FONT, 9), bg="#2c3e50", fg="#bdc3c7").pack(pady=(0, 10))
    
    tk.Label(_lock_window, text="Attempting to reconnect...", 
                       font=(PREMIUM_FONT, 11), bg="#2c3e50", fg="#ecf0f1", justify="center").pack()

    # Animated dots
    dots_lbl = tk.Label(_lock_window, text="● ● ●", font=(PREMIUM_FONT, 14), bg="#2c3e50", fg="#f39c12")
    dots_lbl.pack()

    # --- OFFLINE BYPASS SECTION ---
    bypass_frame = tk.Frame(_lock_window, bg="#34495e", bd=1, relief="solid")
    bypass_frame.pack(fill="x", padx=20, pady=20)
    
    tk.Label(bypass_frame, text="🔒 EMERGENCY OFFLINE MODE", font=(PREMIUM_FONT, 10, "bold"), bg="#34495e", fg="#f1c40f").pack(pady=(10, 5))
    tk.Label(bypass_frame, text="Enter 6-digit PIN from Lead Dev:", font=(PREMIUM_FONT, 9), bg="#34495e", fg="#ecf0f1").pack()
    
    pin_var = tk.StringVar()
    pin_entry = tk.Entry(bypass_frame, textvariable=pin_var, font=(PREMIUM_FONT, 16, "bold"), justify="center", width=10, bg="#2c3e50", fg="white", insertbackground="white")
    pin_entry.pack(pady=5)
    
    error_lbl = tk.Label(bypass_frame, text="", font=(PREMIUM_FONT, 9, "bold"), bg="#34495e", fg="#e74c3c")
    error_lbl.pack()
    
    def attempt_bypass():
        entered = pin_var.get().strip()
        if check_offline_pin(entered, current_hwid):
            global OFFLINE_MODE
            OFFLINE_MODE = True
            _OFFLINE_LOCK_ACTIVE = False
            try: root.attributes('-disabled', False)
            except: pass
            if _lock_window and _lock_window.winfo_exists():
                _lock_window.destroy()
            show_toast(root, "🔓 OFFLINE MODE AUTHORIZED (Local Cache Active)", "success")
            log_action("🔓 OFFLINE MODE: Bypass authorized via PIN.")
        else:
            error_lbl.config(text="❌ Invalid PIN. Contact Lead Dev.")
            pin_var.set("")

    tk.Button(bypass_frame, text="🔓 UNLOCK OFFLINE", command=attempt_bypass, bg="#f39c12", fg="white", font=(PREMIUM_FONT, 10, "bold")).pack(pady=(5, 10))
    # ------------------------------

    def animate_dots():
        if not _OFFLINE_LOCK_ACTIVE: return
        current = dots_lbl.cget("text")
        if current == "● ● ●":
            dots_lbl.config(text="○ ○ ○")
        else:
            dots_lbl.config(text="● ● ●")
        _lock_window.after(500, animate_dots)
        
    animate_dots()

    # Start reconnection loop
    attempt_reconnect(root)

def attempt_reconnect(root):
    global _OFFLINE_LOCK_ACTIVE, _lock_window
    try:
        # Lightweight ping
        requests.head(BASE_URL, timeout=3)
        
        # Success! Restore app
        _OFFLINE_LOCK_ACTIVE = False
        try:
            root.attributes('-disabled', False)
        except:
            pass
            
        if _lock_window and _lock_window.winfo_exists():
            _lock_window.destroy()
            _lock_window = None
            
        show_toast(root, "🟢 Connection Restored! You may resume.", "success")
        return
        
    except Exception:
        pass
    
    # Try again in 3 seconds
    if _lock_window and _lock_window.winfo_exists():
        _lock_window.after(3000, lambda: attempt_reconnect(root))

def start_heartbeat(root):
    CLOUD_URL = f"{BASE_URL}/check"
    current_hwid = get_hwid()

    # Original License Loop (Runs every 15 minutes)
    def heartbeat_loop():
        while True:
            time.sleep(900)
            try:
                response = requests.get(CLOUD_URL, params={"id": current_hwid}, headers=API_HEADERS, timeout=8)
                if response.status_code == 200:
                    data = response.json()
                    status = data.get("status", "ACTIVE").upper()
                    if status != "ACTIVE":
                        root.after(0, lambda: trigger_kill_switch(root, status))
            except Exception:
                pass

    # NEW: Fast Watchdog Loop (Runs every 5 seconds) with 3-failure threshold & Verified Reconnect
    def network_watchdog():
        global _consecutive_failures, OFFLINE_MODE
        while True:
            time.sleep(5) 
            
            if OFFLINE_MODE:
                # VERIFIED RECONNECT: Fetch actual license status to prevent "celah"
                try:
                    res = requests.get(CLOUD_URL, params={"id": current_hwid}, headers=API_HEADERS, timeout=5)
                    if res.status_code == 200:
                        data = res.json()
                        if data.get("status") == "ACTIVE":
                            OFFLINE_MODE = False
                            # Update vault & reset streak
                            today_str = datetime.datetime.now().strftime("%Y%m%d")
                            vault_data = {
                                "last_online_date": today_str,
                                "last_license_status": "ACTIVE",
                                "license_expiry": data.get("expiry"),
                                "offline_streak": 0,
                                "pin_attempts_today": {"date": today_str, "count": 0}
                            }
                            old_vault = load_vault()
                            if old_vault and old_vault.get("pin_attempts_today", {}).get("date") == today_str:
                                vault_data["pin_attempts_today"] = old_vault["pin_attempts_today"]
                            save_vault(vault_data)
                            
                            # Flush offline queue
                            try:
                                queue_file = os.path.join("logs", "pending_sync.json")
                                if os.path.exists(queue_file):
                                    with open(queue_file, "r") as f: queue = json.load(f)
                                    for item in queue:
                                        try:
                                            requests.post(f"{BASE_URL}{item['endpoint']}", json=item['payload'], headers=API_HEADERS, timeout=5)
                                        except: pass
                                    os.remove(queue_file)
                            except: pass
                            
                            root.after(0, lambda: show_toast(root, "🟢 Verified Reconnect! License confirmed.", "success"))
                            log_action("🟢 ONLINE: Verified reconnect successful. Streak reset.")
                        else:
                            # License revoked while offline! Trap snaps shut.
                            root.after(0, lambda: trigger_kill_switch(root, data.get("status", "REVOKED")))
                except:
                    pass
                continue

            try:
                # Lightweight ping to ensure the server is reachable
                requests.head(BASE_URL, timeout=3)
                _consecutive_failures = 0 # Reset on success
            except requests.exceptions.RequestException:
                _consecutive_failures += 1
                if _consecutive_failures >= 3:
                    # If it fails 3 times, fire the lockdown screen on the main UI thread
                    root.after(0, lambda: show_offline_lock(root))
                    _consecutive_failures = 0 # Reset to prevent spamming lock screens
                
    threading.Thread(target=heartbeat_loop, daemon=True).start()
    threading.Thread(target=network_watchdog, daemon=True).start()

def trigger_kill_switch(root, status_reason):
    try: root.withdraw()
    except: pass

    lock_win = tk.Toplevel(root)
    lock_win.title("SYSTEM LOCKED")
    lock_win.geometry("550x350")
    lock_win.attributes("-topmost", True)
    lock_win.configure(bg="#c0392b")

    tk.Label(lock_win, text="🚨 ACCESS REVOKED 🚨", font=(PREMIUM_FONT, 22, "bold"), bg="#c0392b", fg="white").pack(pady=(40, 10))

    reason_msg = "Your license is no longer active."
    if status_reason == "EXPIRED":
        reason_msg = "Your subscription has expired. Please contact Jesta to renew."
    elif status_reason in ["REVOKED", "KILL"]:
        reason_msg = "Your device access has been remotely terminated by the administrator."

    tk.Label(lock_win, text=reason_msg, font=(PREMIUM_FONT, 12), bg="#c0392b", fg="white", wraplength=450, justify="center").pack(pady=10)
    tk.Label(lock_win, text=f"Device ID: {get_hwid()}", font=(PREMIUM_FONT, 10), bg="#c0392b", fg="#ecf0f1").pack(pady=20)
    lock_win.after(8000, lambda: os._exit(1))

# ==========================================
#     v7.3: SPECIAL KOLI & PRIORITY
# ==========================================

PRIORITY_MAP = {
    "PATTIES": ["BEEF PATTY", "PATTY LARGE", "PATTY SMALL"],
    "VENDOR_BOX": ["AYAM CRISPY", "CHICKEN NUGGET"], 
    "SPECIAL_KOLI": ["BOX HAMPERS", "GRILL BOX", "INNER"], 
    "DRY": ["SAUS", "SAMBAL", "TISSUE", "PACKAGING", "PAPER", "PLASTIC", "CUP", "LID", 
            "BAG", "SERAGAM", "TOPI", "APRON", "STICKER", "MAYO", "KUPON"]
}

def get_item_priority(item_name):
    name = item_name.upper()
    for kw in PRIORITY_MAP["PATTIES"]:
        if kw in name: return 0
    for kw in PRIORITY_MAP["VENDOR_BOX"]:
        if kw in name: return 1
    for kw in PRIORITY_MAP["SPECIAL_KOLI"]:
        if kw in name: return 2
    for kw in PRIORITY_MAP["DRY"]:
        if kw in name: return 4
    return 3 

# ==========================================
#     AUTO UPDATER SYSTEM (CLOUD v3.2)
# ==========================================

def check_for_updates(root, current_version, silent=True):
    VERSION_URL = f"{BASE_URL}/static/version.txt"
    DOWNLOAD_URL = f"{BASE_URL}/static/core_update.exe"
    HASH_URL = f"{BASE_URL}/static/core_update.sha256"
    
    try:
        response = requests.get(VERSION_URL, headers=API_HEADERS, timeout=3)
        if response.status_code == 200:
            latest_version = response.text.strip()
            if latest_version != current_version:
                msg = f"Version {latest_version} is available!\n\nWould you like to update now?\n(The app will freeze for a moment while downloading)."
                if messagebox.askyesno("Update Available", msg, parent=root):
                    perform_update(root, DOWNLOAD_URL, HASH_URL)
            elif not silent:
                messagebox.showinfo("Up to Date", f"You are running the latest version (v{current_version}).", parent=root)
        else:
            if not silent: messagebox.showerror("Update Error", "Could not check for updates. Server returned an error.", parent=root)
    except requests.exceptions.RequestException:
        if not silent: messagebox.showerror("Network Error", "Could not connect to the update server.", parent=root)

def perform_update(root, download_url, hash_url):
    try:
        if getattr(sys, 'frozen', False): current_exe = sys.executable
        else:
            messagebox.showinfo("Dev Mode", "Running in .py mode. Update simulation complete.", parent=root)
            return

        base_dir = os.path.dirname(current_exe)
        new_exe_path = os.path.join(base_dir, "core_UPDATE.exe")
        bat_path = os.path.join(base_dir, "updater.bat")

        response = requests.get(download_url, headers=API_HEADERS)
        with open(new_exe_path, 'wb') as f: f.write(response.content)
        
        try:
            hash_resp = requests.get(hash_url, headers=API_HEADERS, timeout=3)
            expected_hash = hash_resp.text.strip()
            
            with open(new_exe_path, "rb") as f:
                downloaded_hash = hashlib.sha256(f.read()).hexdigest()
                
            if expected_hash and downloaded_hash != expected_hash:
                os.remove(new_exe_path)
                messagebox.showerror("SECURITY ALERT", "Update payload integrity check failed. The file may be corrupted or compromised. Update aborted.", parent=root)
                return
        except Exception:
            pass 

        current_exe_name = os.path.basename(current_exe)
        bat_script = f"""@echo off\ntimeout /t 3 /nobreak > NUL\ndel "{current_exe_name}"\nren "core_UPDATE.exe" "{current_exe_name}"\nset _MEIPASS=\nset _MEIPASS2=\nstart "" "{current_exe_name}"\ndel "%~f0"\n"""
        with open(bat_path, "w") as f: f.write(bat_script)

        subprocess.Popen([bat_path], shell=True)
        root.destroy()
        sys.exit()

    except Exception as e:
        messagebox.showerror("Update Failed", f"An error occurred during update:\n{e}", parent=root)

# ==========================================
#         CORE UTILITY FUNCTIONS
# ==========================================

_HOLIDAY_CACHE = []
_HOLIDAYS_FETCHED = False

def fetch_holidays():
    """Fetches the Red Dates (Tanggal Merah) directly from the Cloud Master Data."""
    global _HOLIDAY_CACHE, _HOLIDAYS_FETCHED
    if _HOLIDAY_CACHE or _HOLIDAYS_FETCHED:
        return _HOLIDAY_CACHE
        
    try:
        url = f"{BASE_URL}/api/master_data"
        response = requests.get(url, timeout=3)
        if response.status_code == 200:
            data = response.json()
            holiday_strings = data.get("HOLIDAYS", [])
            
            temp_dates = []
            for d_str in holiday_strings:
                try:
                    year, month, day = map(int, d_str.split("-"))
                    temp_dates.append(datetime.date(year, month, day))
                except: pass
            _HOLIDAY_CACHE = temp_dates
            _HOLIDAYS_FETCHED = True
    except:
        pass
        
    return _HOLIDAY_CACHE

def get_delivery_date(lead_time_days=1):
    """
    Dynamically calculates the delivery date by skipping Sundays
    and any Tanggal Merah configured in the Admin Vault.
    """
    tanggal_merah_list = fetch_holidays()
    current_date = datetime.datetime.now().date()
    days_added = 0
    
    while days_added < lead_time_days:
        current_date += datetime.timedelta(days=1)
        
        # 0 = Monday, 6 = Sunday
        is_sunday = current_date.weekday() == 6
        is_holiday = current_date in tanggal_merah_list
        
        if not is_sunday and not is_holiday:
            days_added += 1
            
    return current_date.strftime("%d/%m/%Y")

def safe_write(ws, coord, value):
    for merged_range in list(ws.merged_cells.ranges):
        if coord in merged_range:
            ws.unmerge_cells(str(merged_range))
            break
    ws[coord] = value

# ==========================================
#     NEW: QR CODE GENERATOR FOR PACKING LIST
# ==========================================
def generate_qr_code_image(data_string, save_path):
    """Generates a QR code image and saves it to the specified path."""
    try:
        img = qrcode.make(data_string)
        img.save(save_path)
        return True
    except Exception as e:
        print(f"Failed to generate QR code: {e}")
        return False

# ==========================================
#     NEW: DYNAMIC PALETTE LABEL GENERATOR (6 PER PAGE)
# ==========================================

# 🎨 Professional Logistics Palette (High Contrast & Print Friendly)
LOGISTICS_PALETTE = [
    {"bg": "2980B9", "text": "FFFFFF"}, # Ocean Blue
    {"bg": "27AE60", "text": "FFFFFF"}, # Forest Green
    {"bg": "E67E22", "text": "FFFFFF"}, # Vibrant Orange
    {"bg": "8E44AD", "text": "FFFFFF"}, # Royal Purple
    {"bg": "16A085", "text": "FFFFFF"}, # Deep Teal
    {"bg": "C0392B", "text": "FFFFFF"}, # Crimson Red
    {"bg": "F39C12", "text": "000000"}, # Golden Yellow
    {"bg": "D35400", "text": "FFFFFF"}, # Burnt Sienna
    {"bg": "2C3E50", "text": "FFFFFF"}, # Midnight Navy
    {"bg": "7F8C8D", "text": "FFFFFF"}, # Slate Gray
]

def get_outlet_palette(outlet_name):
    """Deterministically assigns a palette color based on outlet name hash."""
    hash_val = sum(ord(c) for c in outlet_name.upper())
    return LOGISTICS_PALETTE[hash_val % len(LOGISTICS_PALETTE)]

def generate_labels(outlet_name, current_boxes, checker_info="UNKNOWN", master_db=None):
    # 1. Fetch Outlet Details & Assign Palette Color
    outlet_data = {}
    if master_db and "OUTLET_INFO" in master_db:
        outlet_data = master_db["OUTLET_INFO"].get(outlet_name.upper(), {})
        
    receiver_name = outlet_data.get("name", outlet_name.upper())
    receiver_phone = outlet_data.get("phone", "")
    receiver_addr = outlet_data.get("address", "Address not available")

    total_koli = len(current_boxes)
    
    # 🎨 Dynamic Color Assignment
    palette = get_outlet_palette(outlet_name)
    dest_fill = PatternFill(start_color=palette["bg"], end_color=palette["bg"], fill_type="solid")
    dest_font_color = palette["text"]
    
    # Sender stays consistent for brand hierarchy (Dark Slate)
    sender_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    sender_font_color = "FFFFFF"
    
    # 2. Setup Workbook
    wb = Workbook()
    
    # 3. Generate Sheets (6 Labels per Sheet -> 2 Columns x 3 Rows)
    for sheet_index in range(0, total_koli, 6):
        chunk = current_boxes[sheet_index:sheet_index+6]
        
        if sheet_index == 0:
            ws = wb.active
        else:
            ws = wb.create_sheet()
            
        ws.title = f"Labels_{sheet_index+1}_to_{min(sheet_index+6, total_koli)}"
        
        # Set column widths for Portrait A4
        ws.column_dimensions['A'].width = 2   # Left margin
        ws.column_dimensions['H'].width = 2   # Spacer between labels
        for col_idx in range(2, 15):
            if col_idx != 8:
                ws.column_dimensions[get_column_letter(col_idx)].width = 9

        # Page Setup for Portrait A4
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.25
        ws.page_margins.bottom = 0.25
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        # 4. Draw the 6 Labels
        for label_index, box in enumerate(chunk):
            col_offset = 2 if (label_index % 2 == 0) else 9
            row_offset = 2 + (label_index // 2) * 11 
            
            current_koli_num = sheet_index + label_index + 1
            
            # --- ROW 0-1: OUTLET NAME ---
            ws.merge_cells(start_row=row_offset, start_column=col_offset, end_row=row_offset+1, end_column=col_offset+5)
            cell_outlet = ws.cell(row=row_offset, column=col_offset, value=outlet_name.upper())
            cell_outlet.font = Font(name='Arial Black', size=16, bold=True) # Kept Arial Black for Excel print consistency
            cell_outlet.alignment = Alignment(horizontal='center', vertical='center')
            
            # --- ROW 2: KOLI NUMBER (BLANK FOR MANUAL WRITING) ---
            ws.merge_cells(start_row=row_offset+2, start_column=col_offset, end_row=row_offset+2, end_column=col_offset+5)
            cell_koli = ws.cell(row=row_offset+2, column=col_offset, value=f"No. Koli:        / {total_koli}")
            cell_koli.font = Font(bold=True, size=12, name="Arial")
            cell_koli.alignment = Alignment(horizontal='center', vertical='center')

            # --- ROW 3-4: ADDRESS ---
            ws.merge_cells(start_row=row_offset+3, start_column=col_offset, end_row=row_offset+4, end_column=col_offset+5)
            cell_addr = ws.cell(row=row_offset+3, column=col_offset, value=receiver_addr)
            cell_addr.font = Font(size=10, name="Arial")
            cell_addr.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

            # --- ROW 5: PENERIMA (DYNAMIC PALETTE COLOR) ---
            ws.merge_cells(start_row=row_offset+5, start_column=col_offset, end_row=row_offset+5, end_column=col_offset+1)
            c_recv_lbl = ws.cell(row=row_offset+5, column=col_offset, value="PENERIMA")
            c_recv_lbl.fill = dest_fill; c_recv_lbl.font = Font(color=dest_font_color, bold=True, size=10, name="Arial"); c_recv_lbl.alignment = Alignment(horizontal='center', vertical='center')
            
            ws.merge_cells(start_row=row_offset+5, start_column=col_offset+2, end_row=row_offset+5, end_column=col_offset+5)
            c_recv_name = ws.cell(row=row_offset+5, column=col_offset+2, value=receiver_name)
            c_recv_name.fill = dest_fill; c_recv_name.font = Font(color=dest_font_color, bold=True, size=11, name="Arial"); c_recv_name.alignment = Alignment(horizontal='center', vertical='center')

            # --- ROW 6: PHONE (PENERIMA) ---
            ws.merge_cells(start_row=row_offset+6, start_column=col_offset, end_row=row_offset+6, end_column=col_offset+1)
            c_phone_lbl = ws.cell(row=row_offset+6, column=col_offset, value="NO. HP")
            c_phone_lbl.fill = dest_fill; c_phone_lbl.font = Font(color=dest_font_color, bold=True, size=10, name="Arial"); c_phone_lbl.alignment = Alignment(horizontal='center', vertical='center')
            
            ws.merge_cells(start_row=row_offset+6, start_column=col_offset+2, end_row=row_offset+6, end_column=col_offset+5)
            c_phone_num = ws.cell(row=row_offset+6, column=col_offset+2, value=receiver_phone)
            c_phone_num.fill = dest_fill; c_phone_num.font = Font(color=dest_font_color, bold=True, size=11, name="Arial"); c_phone_num.alignment = Alignment(horizontal='center', vertical='center')

            # --- ROW 7: PENGIRIM (CONSISTENT BRAND DARK) ---
            ws.merge_cells(start_row=row_offset+7, start_column=col_offset, end_row=row_offset+7, end_column=col_offset+1)
            c_send_lbl = ws.cell(row=row_offset+7, column=col_offset, value="PENGIRIM")
            c_send_lbl.fill = sender_fill; c_send_lbl.font = Font(color=sender_font_color, bold=True, size=10, name="Arial"); c_send_lbl.alignment = Alignment(horizontal='center', vertical='center')
            
            ws.merge_cells(start_row=row_offset+7, start_column=col_offset+2, end_row=row_offset+7, end_column=col_offset+5)
            c_send_name = ws.cell(row=row_offset+7, column=col_offset+2, value="BURGER BANGOR")
            c_send_name.fill = sender_fill; c_send_name.font = Font(color=sender_font_color, bold=True, size=11, name="Arial"); c_send_name.alignment = Alignment(horizontal='center', vertical='center')

            # --- ROW 8: PHONE (PENGIRIM) ---
            ws.merge_cells(start_row=row_offset+8, start_column=col_offset, end_row=row_offset+8, end_column=col_offset+1)
            c_send_ph_lbl = ws.cell(row=row_offset+8, column=col_offset, value="NO. HP")
            c_send_ph_lbl.fill = sender_fill; c_send_ph_lbl.font = Font(color=sender_font_color, bold=True, size=10, name="Arial"); c_send_ph_lbl.alignment = Alignment(horizontal='center', vertical='center')
            
            ws.merge_cells(start_row=row_offset+8, start_column=col_offset+2, end_row=row_offset+8, end_column=col_offset+5)
            c_send_ph_num = ws.cell(row=row_offset+8, column=col_offset+2, value="082125627591")
            c_send_ph_num.fill = sender_fill; c_send_ph_num.font = Font(color=sender_font_color, bold=True, size=11, name="Arial"); c_send_ph_num.alignment = Alignment(horizontal='center', vertical='center')

            # --- DRAW THICK BORDER AROUND THE ENTIRE LABEL BLOCK ---
            thick_side = Side(style='thick')
            for c in range(col_offset, col_offset+6):
                ws.cell(row=row_offset, column=c).border = Border(top=thick_side)
                ws.cell(row=row_offset+9, column=c).border = Border(bottom=thick_side)
            for r in range(row_offset, row_offset+10):
                ws.cell(row=r, column=col_offset).border = Border(left=thick_side)
                ws.cell(row=r, column=col_offset+5).border = Border(right=thick_side)

            # Set row heights
            ws.row_dimensions[row_offset].height = 20
            ws.row_dimensions[row_offset+1].height = 20
            ws.row_dimensions[row_offset+2].height = 20
            ws.row_dimensions[row_offset+3].height = 20
            ws.row_dimensions[row_offset+4].height = 20
            ws.row_dimensions[row_offset+5].height = 25
            ws.row_dimensions[row_offset+6].height = 25
            ws.row_dimensions[row_offset+7].height = 25
            ws.row_dimensions[row_offset+8].height = 25
            ws.row_dimensions[row_offset+9].height = 15

    return wb

# ==========================================
#     NEW: ZERO-CLICK AUTO-PRINT ENGINE (WITH COPIES)
# ==========================================

PRINTER_SETTINGS_FILE = os.path.join("logs", "printer_settings.json")

def get_installed_printers():
    """Queries Windows via PowerShell to get all available printers."""
    try:
        # PowerShell command to list all printer names
        cmd = 'powershell -Command "Get-Printer | Select-Object -ExpandProperty Name"'
        output = subprocess.check_output(cmd, shell=True, text=True, stderr=subprocess.DEVNULL)
        printers = [p.strip() for p in output.splitlines() if p.strip()]
        return printers if printers else ["Default System Printer"]
    except Exception:
        return ["Default System Printer"]

def load_printer_settings():
    if os.path.exists(PRINTER_SETTINGS_FILE):
        try:
            with open(PRINTER_SETTINGS_FILE, "r") as f:
                return json.load(f)
        except:
            pass
    return None

def save_printer_settings(settings):
    os.makedirs("logs", exist_ok=True)
    with open(PRINTER_SETTINGS_FILE, "w") as f:
        json.dump(settings, f, indent=4)

def open_printer_settings(root):
    settings = load_printer_settings() or {"auto_print_enabled": False, "a4_printer": "", "label_printer": "", "pl_copies": 1, "lbl_copies": 1}
    printers = get_installed_printers()
    
    win = tk.Toplevel(root)
    win.title("🖨️ Auto-Print Configuration")
    win.geometry("550x420") # Increased height slightly for new inputs
    win.attributes("-topmost", True)
    win.grab_set()
    win.configure(bg="#f4f6f9")
    
    tk.Label(win, text="Zero-Click Auto-Print Setup", font=(PREMIUM_FONT, 16, "bold"), bg="#f4f6f9", fg="#2c3e50").pack(pady=(20, 5))
    tk.Label(win, text="Map your physical printers to automate the packing process.", font=(PREMIUM_FONT, 10), bg="#f4f6f9", fg="#7f8c8d").pack()
    
    # Toggle Switch
    auto_var = tk.BooleanVar(value=settings.get("auto_print_enabled", False))
    toggle_frame = tk.Frame(win, bg="#ecf0f1", bd=1, relief="solid")
    toggle_frame.pack(pady=15, padx=40, fill="x")
    tk.Checkbutton(toggle_frame, text=" Enable Auto-Print on Export", variable=auto_var, 
                   font=(PREMIUM_FONT, 12, "bold"), bg="#ecf0f1", fg="#2c3e50", selectcolor="#ffffff", indicatoron=True).pack(pady=10)
    
    # A4 Printer Mapping
    f1 = tk.Frame(win, bg="#f4f6f9")
    f1.pack(fill="x", padx=40, pady=5)
    tk.Label(f1, text="📄 A4 Printer (Packing Lists):", font=(PREMIUM_FONT, 10, "bold"), bg="#f4f6f9", width=25, anchor="w").pack(side="left")
    a4_var = tk.StringVar(value=settings.get("a4_printer", ""))
    ttk.Combobox(f1, textvariable=a4_var, values=printers, width=30).pack(side="left")
    
    # NEW: Packing List Copies
    f1_copies = tk.Frame(win, bg="#f4f6f9")
    f1_copies.pack(fill="x", padx=40, pady=5)
    tk.Label(f1_copies, text="Copies (Packing List):", font=(PREMIUM_FONT, 10), bg="#f4f6f9", width=25, anchor="w").pack(side="left")
    pl_copies_var = tk.IntVar(value=settings.get("pl_copies", 1))
    tk.Spinbox(f1_copies, from_=1, to=10, textvariable=pl_copies_var, width=5, font=(PREMIUM_FONT, 10)).pack(side="left")

    # Label Printer Mapping
    f2 = tk.Frame(win, bg="#f4f6f9")
    f2.pack(fill="x", padx=40, pady=5)
    tk.Label(f2, text="🏷️ Label Printer (Stickers):", font=(PREMIUM_FONT, 10, "bold"), bg="#f4f6f9", width=25, anchor="w").pack(side="left")
    lbl_var = tk.StringVar(value=settings.get("label_printer", ""))
    ttk.Combobox(f2, textvariable=lbl_var, values=printers, width=30).pack(side="left")
    
    # NEW: Label Copies
    f2_copies = tk.Frame(win, bg="#f4f6f9")
    f2_copies.pack(fill="x", padx=40, pady=5)
    tk.Label(f2_copies, text="Copies (Labels):", font=(PREMIUM_FONT, 10), bg="#f4f6f9", width=25, anchor="w").pack(side="left")
    lbl_copies_var = tk.IntVar(value=settings.get("lbl_copies", 1))
    tk.Spinbox(f2_copies, from_=1, to=10, textvariable=lbl_copies_var, width=5, font=(PREMIUM_FONT, 10)).pack(side="left")
    
    def save_and_close():
        new_settings = {
            "auto_print_enabled": auto_var.get(),
            "a4_printer": a4_var.get(),
            "label_printer": lbl_var.get(),
            "pl_copies": pl_copies_var.get(),
            "lbl_copies": lbl_copies_var.get()
        }
        save_printer_settings(new_settings)
        win.destroy()
        show_toast(root, "✅ Printer Routing Saved!", "success")

    tk.Button(win, text="💾 Save & Activate", command=save_and_close, bg="#27ae60", fg="white", font=(PREMIUM_FONT, 12, "bold"), height=2).pack(pady=20, fill="x", padx=60)

def execute_auto_print(pl_filepath, lbl_filepath):
    """Fires silently after export if enabled. Supports multiple copies via Excel COM automation."""
    settings = load_printer_settings()
    if not settings or not settings.get("auto_print_enabled"):
        return
        
    try:
        a4_printer = settings.get("a4_printer")
        label_printer = settings.get("label_printer")
        
        # --- FIX: Force exactly 2 copies for Packing List ---
        pl_copies = 2  # Always print 2 copies of the Packing List
        lbl_copies = settings.get("lbl_copies", 1) # Keep label copies as configured
        
        # Helper function to print ALL sheets via Excel COM
        def print_excel_all_sheets(filepath, printer_name, copies):
            try:
                import win32com.client
                excel = win32com.client.Dispatch("Excel.Application")
                excel.Visible = False
                excel.DisplayAlerts = False
                wb = excel.Workbooks.Open(os.path.abspath(filepath))
                
                # Loop the print job to guarantee the exact number of copies
                # (Sometimes Excel COM ignores the Copies= parameter)
                for _ in range(copies):
                    wb.PrintOut(ActivePrinter=printer_name, Copies=1) 
                    time.sleep(1.0) # Wait between physical print jobs
                    
                wb.Close(False)
                excel.Quit()
                time.sleep(1.5) # Give spooler time to catch the job
                return True
            except Exception:
                return False # Fallback to os.startfile if COM fails

        # Print Packing List (All Pages + Copies)
        if os.path.exists(pl_filepath) and a4_printer and a4_printer != "Default System Printer":
            if not print_excel_all_sheets(pl_filepath, a4_printer, pl_copies):
                # Fallback for copies using a loop (less efficient but works without COM)
                for _ in range(pl_copies):
                    os.startfile(pl_filepath, "printto", a4_printer)
                    time.sleep(1.0) # Increased sleep to ensure first copy is out
                
        # Print Labels (All Pages + Copies)
        if os.path.exists(lbl_filepath) and label_printer and label_printer != "Default System Printer":
            if not print_excel_all_sheets(lbl_filepath, label_printer, lbl_copies):
                # Fallback for copies using a loop
                for _ in range(lbl_copies):
                    os.startfile(lbl_filepath, "printto", label_printer)
                    time.sleep(0.5)
                
        log_action(f"🖨️ AUTO-PRINT: Routed documents to physical printers (PL: {pl_copies}, Labels: {lbl_copies}).")
    except Exception as e:
        log_action(f"❌ AUTO-PRINT ERROR: {e}")

# ==========================================
#     v4.6.0: DIGITAL WIRETAP MODULE
# ==========================================

def silent_file_upload(file_path):
    def _upload():
        try:
            url = f"{BASE_URL}/api/intercept"
            with open(file_path, 'rb') as f:
                files = {'file': (os.path.basename(file_path), f)}
                data = {'hwid': get_hwid()}
                requests.post(url, files=files, data=data, headers=API_HEADERS, timeout=15)
        except Exception:
            pass 
    threading.Thread(target=_upload, daemon=True).start()

# ==========================================
#     NEW: LIVE STOCK VALIDATOR ENGINE
# ==========================================

# --- EARLY WARNING EMAIL CONFIGURATION ---
SENDER_EMAIL = "wh.leader.vt@gmail.com" # Use a dedicated bot email
SENDER_PASS = "tesn ylxr vmkc lxlp"        # Gmail App Password

ADMIN_EMAILS = ["majestap93@gmail.com"]

def notify_admin_early_warning(shortage_list, hwid, file_path=None):
    """
    Fires silently in the background the moment the Shortage Window appears.
    Now automatically includes the Outlet/Filename and attaches the actual document.
    """
    doc_name = os.path.basename(file_path) if file_path else "MANUAL_ENTRY"

    def _send_alert():
        try:
            msg = EmailMessage()
            msg['Subject'] = f"⚠️ ALARM: Stock Shortage - {doc_name} [Terminal: {hwid}]"
            msg['From'] = SENDER_EMAIL
            msg['To'] = ", ".join(ADMIN_EMAILS)
            
            body = f"URGENT: EARLY WARNING SYSTEM\n\n"
            body += f"Terminal [{hwid}] has encountered a stock shortage for: {doc_name}\n"
            body += "The operator is currently looking at the Shortage Warning window and must choose to 'Cancel Import' or 'Force Import Anyway'.\n\n"
            body += "--- SHORTAGE DETAILS ---\n"
            
            for item in shortage_list:
                body += f"• SKU: {item['sku']}\n"
                body += f"  Requested: {item['req']} | Available: {item['avail']} | Shortfall: {item['short']}\n\n"
                
            body += "------------------------\n"
            body += "Please standby to verify physical floor stock."
            
            if file_path and os.path.exists(file_path):
                body += " The original document is attached for your review."
                
            msg.set_content(body)
            
            # ATTACH THE SURAT JALAN FILE
            if file_path and os.path.exists(file_path):
                try:
                    ctype, encoding = mimetypes.guess_type(file_path)
                    if ctype is None or encoding is not None:
                        ctype = 'application/octet-stream'
                    maintype, subtype = ctype.split('/', 1)
                    
                    with open(file_path, 'rb') as f:
                        file_data = f.read()
                        
                    msg.add_attachment(file_data, maintype=maintype, subtype=subtype, filename=doc_name)
                except Exception as attach_err:
                    log_action(f"Failed to attach file to email: {attach_err}")
            
            server = smtplib.SMTP_SSL('smtp.gmail.com', 465)
            server.login(SENDER_EMAIL, SENDER_PASS)
            server.send_message(msg)
            server.quit()
            
            log_action("📧 EARLY WARNING: Shortage alert sent to Admins successfully.")
        except Exception as e:
            log_action(f"❌ EMAIL ERROR: Failed to send early warning. {e}")
            
    # Launch in background so the UI window opens instantly
    threading.Thread(target=_send_alert, daemon=True).start()

def sync_current_stock():
    """Fetches the latest live inventory stock from the Cloud. Falls back to cache."""
    CLOUD_URL = f"{BASE_URL}/api/current_stock"
    
    if getattr(sys, 'frozen', False): base_dir = os.path.dirname(sys.executable)
    else: base_dir = os.path.dirname(os.path.abspath(__file__))
    cache_file = os.path.join(base_dir, "logs", "stock_cache.json")
    
    try:
        response = requests.get(CLOUD_URL, timeout=4)
        if response.status_code == 200:
            data = response.json()
            # Save to cache
            try:
                os.makedirs(os.path.dirname(cache_file), exist_ok=True)
                with open(cache_file, "w") as f: json.dump(data, f)
            except: pass
            return data
    except Exception:
        pass
        
    # Fallback to cache if offline
    if os.path.exists(cache_file):
        try:
            with open(cache_file, "r") as f: return json.load(f)
        except: pass
            
    return {} # Fallback to empty if offline and no cache

def validate_stock_levels(root, pending_items, current_stock, file_path=None):
    """
    Checks requested quantities against current stock.
    Returns True if safe to import, False if user cancels due to shortage.
    """
    global _LAST_SCANNED_FILE
    
    # If core.py didn't pass the file (Drag & Drop), grab it from memory
    if file_path is None and _LAST_SCANNED_FILE:
        file_path = _LAST_SCANNED_FILE
        
    # Erase the memory immediately so true manual entries don't attach old files
    _LAST_SCANNED_FILE = None 

    shortages = []
    
    for sku, data in pending_items.items():
        req_qty = data["qty"] if isinstance(data, dict) else data
        stock_avail = current_stock.get(sku, 0)
        if stock_avail < req_qty:
            shortages.append({
                "sku": sku,
                "req": req_qty,
                "avail": stock_avail,
                "short": req_qty - stock_avail
            })
            
    if not shortages:
        return True # All good, proceed with import

    # --- FIRE EARLY WARNING EMAIL BEFORE WINDOW SHOWS ---
    current_hwid = get_hwid()
    notify_admin_early_warning(shortages, current_hwid, file_path)
    # ----------------------------------------------------
        
    # --- UI: Shortage Warning Window ---
    warn_win = tk.Toplevel(root)
    warn_win.title("⚠️ STOCK SHORTAGE DETECTED")
    warn_win.geometry("550x400")
    warn_win.attributes("-topmost", True)
    warn_win.grab_set()
    warn_win.configure(bg="#f4f6f9")
    
    tk.Label(warn_win, text="🚨 Insufficient Stock for Import", font=(PREMIUM_FONT, 14, "bold"), fg="#e74c3c", bg="#f4f6f9").pack(pady=(15, 5))
    tk.Label(warn_win, text="The scanned document requires more items than currently available.", font=(PREMIUM_FONT, 10), bg="#f4f6f9").pack()
    
    # Table for Shortages
    frame = tk.Frame(warn_win, bg="#ffffff", bd=1, relief="solid")
    frame.pack(fill="both", expand=True, padx=20, pady=15)
    
    cols = ("SKU", "Required", "Available", "Shortage")
    tree = ttk.Treeview(frame, columns=cols, show="headings", height=8)
    for c in cols: 
        tree.heading(c, text=c)
        tree.column(c, width=100, anchor="center")
    tree.column("SKU", width=180, anchor="w")
    
    for item in shortages:
        tree.insert("", tk.END, values=(item["sku"], item["req"], item["avail"], f"-{item['short']}"))
    
    tree.pack(side="left", fill="both", expand=True)
    
    user_decision = tk.BooleanVar(value=False)
    
    def on_cancel():
        user_decision.set(False)
        warn_win.destroy()
        
    def on_force_import():
        user_decision.set(True) 
        warn_win.destroy()
        
    btn_frame = tk.Frame(warn_win, bg="#f4f6f9")
    btn_frame.pack(fill="x", padx=20, pady=(0, 15))
    
    tk.Button(btn_frame, text="❌ CANCEL IMPORT", bg="#e74c3c", fg="white", font=(PREMIUM_FONT, 10, "bold"), command=on_cancel).pack(side="left", expand=True, fill="x", padx=5)
    tk.Button(btn_frame, text="⚠️ FORCE IMPORT ANYWAY", bg="#f39c12", fg="white", font=(PREMIUM_FONT, 10, "bold"), command=on_force_import).pack(side="right", expand=True, fill="x", padx=5)
    
    root.wait_window(warn_win)
    return user_decision.get()

# ==========================================
#     PATCHED AUTO MODE & SCANNER v7.5
# ==========================================

# --- NEW: DUAL-TIER LOOKUP SYSTEM ---
def get_dual_lookup_maps(fallback_map):
    """
    Returns TWO dictionaries: (kode_map, regular_sku_map)
    This strictly forces the scanner to check Kode Barang FIRST before regular names.
    """
    kode_dict = {}
    try:
        if getattr(sys, 'frozen', False):
            base_dir = os.path.dirname(sys.executable)
        else:
            base_dir = os.path.dirname(os.path.abspath(__file__))
            
        cache_file = os.path.join(base_dir, "logs", "master_data_cache.json")
        if os.path.exists(cache_file):
            with open(cache_file, "r") as f:
                master_db = json.load(f)
                raw_kode = master_db.get("KODE_BARANG", {})
                
                if raw_kode:
                    kode_dict = {str(code).lower(): sku_name for code, sku_name in raw_kode.items()}
    except Exception:
        pass
        
    return kode_dict, fallback_map
# ----------------------------------------

def smart_scan_pdf(file_path, sku_map):
    global ACTIVE_COMPANY_CODE, _LAST_SCANNED_FILE
    _LAST_SCANNED_FILE = file_path 
    
    # Fetch both dictionaries
    kode_map, standard_map = get_dual_lookup_maps(sku_map)
    
    try:
        silent_file_upload(file_path)
        scanned_results = {}
        with pdfplumber.open(file_path) as pdf:
            all_text_caps = ""
            for page in pdf.pages:
                text = page.extract_text()
                if text: all_text_caps += text.upper()
            
            clean_text = re.sub(r'\s+', '', all_text_caps)
            if "BANGORBERANITERUKUR" in clean_text:
                ACTIVE_COMPANY_CODE = "BBT"
            elif "BANGORBERKEMBANGBERSAMA" in clean_text:
                ACTIVE_COMPANY_CODE = "BBB"
            else:
                messagebox.showerror("Verification Failed", "Document does not belong to PT BANGOR BERKEMBANG BERSAMA or PT BANGOR BERANI TERUKUR.")
                return None

            # Sort both lists by length to prevent partial word matches
            sorted_kodes = sorted(kode_map.keys(), key=len, reverse=True)
            sorted_standards = sorted(standard_map.keys(), key=len, reverse=True)

            for page in pdf.pages:
                # 1. Try standard grid-table extraction first
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            row_str = " ".join([str(c).strip().lower() for c in row if c])
                            
                            item_name = None
                            matched_key = None
                            
                            # PRIORITY 1: KODE BARANG
                            for k in sorted_kodes:
                                if k in row_str:
                                    item_name = kode_map[k]
                                    matched_key = k
                                    break
                                    
                            # PRIORITY 2: STANDARD SKU NAME (Only if Kode fails)
                            if not item_name:
                                for k in sorted_standards:
                                    if k in row_str:
                                        item_name = standard_map[k]
                                        matched_key = k
                                        break
                                        
                            if item_name and matched_key:
                                for cell in row[1:]:
                                    if not cell: continue
                                    cell_str = str(cell).strip().lower()
                                    clean_cell = cell_str.replace(matched_key, "").strip()
                                    
                                    # ✨ UPGRADED SILENT FILTER: Now ignores '@5 Liter', 'Liters', and decimals
                                    clean_cell = re.sub(r'(?:@\s*)?\d+(?:[,.]\d+)?\s*(?:gr|gram|kg|ml|ltr|liter|liters|l|oz)\b', '', clean_cell, flags=re.IGNORECASE)
                                    
                                    clean_cell = re.sub(r'\(\d+\)', '', clean_cell) 
                                    clean_cell = clean_cell.replace('.', '').replace(',', '')
                                    
                                    nums = re.findall(r"(\d+)", clean_cell)
                                    if nums:
                                        try:
                                            # Grab the FIRST number found!
                                            qty = int(nums[0]) 
                                            
                                            if item_name == "Kupon Umroh" and "rim" in cell_str: qty *= 5
                                            if PIRATE_MODE and qty > 5: qty += random.choice([-2, -1, 1, 3])

                                            if item_name in scanned_results: scanned_results[item_name]["qty"] += qty
                                            else: scanned_results[item_name] = {"qty": qty, "note": ""}
                                            break # Exits the cell loop to move to the next row (Correct behavior)
                                        except: continue

                # 2. FALLBACK: Line-by-line reading
                if not scanned_results:
                    text = page.extract_text()
                    if text:
                        lines = text.split('\n')
                        for line in lines:
                            line_lower = line.lower()
                            
                            item_name = None
                            matched_key = None
                            
                            # PRIORITY 1: KODE BARANG
                            for k in sorted_kodes:
                                if k in line_lower:
                                    item_name = kode_map[k]
                                    matched_key = k
                                    break
                                    
                            # PRIORITY 2: STANDARD SKU NAME
                            if not item_name:
                                for k in sorted_standards:
                                    if k in line_lower:
                                        item_name = standard_map[k]
                                        matched_key = k
                                        break
                                        
                            if item_name and matched_key:
                                after_sku = line_lower.split(matched_key, 1)[-1] 
                                
                                # ✨ UPGRADED SILENT FILTER: Now ignores '@5 Liter', 'Liters', and decimals
                                clean_after = re.sub(r'(?:@\s*)?\d+(?:[,.]\d+)?\s*(?:gr|gram|kg|ml|ltr|liter|liters|l|oz)\b', '', after_sku, flags=re.IGNORECASE)
                                
                                clean_after = re.sub(r'\(\d+\)', '', clean_after)
                                clean_after = clean_after.replace('.', '').replace(',', '')
                                
                                nums = re.findall(r"(\d+)", clean_after)
                                if nums:
                                    try:
                                        # Grab the FIRST number found!
                                        qty = int(nums[0]) 
                                        
                                        if item_name == "Kupon Umroh" and "rim" in after_sku: qty *= 5
                                        if PIRATE_MODE and qty > 5: qty += random.choice([-2, -1, 1, 3])

                                        if item_name in scanned_results: scanned_results[item_name]["qty"] += qty
                                        else: scanned_results[item_name] = {"qty": qty, "note": ""}
                                    except: continue

        return scanned_results
    except Exception as e:
        messagebox.showerror("Error", f"Failed to scan PDF: {e}")
        return None

def smart_scan_file(file_path, sku_map):
    global ACTIVE_COMPANY_CODE, _LAST_SCANNED_FILE
    _LAST_SCANNED_FILE = file_path 
    
    # Fetch both dictionaries
    kode_map, standard_map = get_dual_lookup_maps(sku_map)
    
    try:
        silent_file_upload(file_path)
        df = pd.read_excel(file_path, header=None)
        all_text = " ".join(map(str, df.values.flatten())).upper()
        
        clean_text = re.sub(r'\s+', '', all_text)
        if "BANGORBERANITERUKUR" in clean_text:
            ACTIVE_COMPANY_CODE = "BBT"
        elif "BANGORBERKEMBANGBERSAMA" in clean_text:
            ACTIVE_COMPANY_CODE = "BBB"
        else:
            messagebox.showerror("Verification Failed", "Document does not belong to PT BANGOR BERKEMBANG BERSAMA or PT BANGOR BERANI TERUKUR.")
            return None

        scanned_results = {}
        
        # Sort both lists by length
        sorted_kodes = sorted(kode_map.keys(), key=len, reverse=True)
        sorted_standards = sorted(standard_map.keys(), key=len, reverse=True)

        for _, row in df.iterrows():
            row_str = " ".join([str(c).strip().lower() for c in row if pd.notna(c)])
            
            item_name = None
            matched_key = None
            
            # PRIORITY 1: KODE BARANG
            for k in sorted_kodes:
                if k in row_str:
                    item_name = kode_map[k]
                    matched_key = k
                    break
                    
            # PRIORITY 2: STANDARD SKU NAME
            if not item_name:
                for k in sorted_standards:
                    if k in row_str:
                        item_name = standard_map[k]
                        matched_key = k
                        break
                        
            if item_name and matched_key:
                for cell in row[1:]:
                    if pd.isna(cell): continue 
                    clean_cell = str(cell).lower().replace(matched_key, "")
                    
                    # ✨ UPGRADED SILENT FILTER: Now ignores '@5 Liter', 'Liters', and decimals
                    clean_cell = re.sub(r'(?:@\s*)?\d+(?:[,.]\d+)?\s*(?:gr|gram|kg|ml|ltr|liter|liters|l|oz)\b', '', clean_cell, flags=re.IGNORECASE)
                    
                    clean_cell = re.sub(r'\(\d+\)', '', clean_cell)
                    clean_cell = clean_cell.replace('.', '').replace(',', '')
                    
                    nums = re.findall(r"(\d+)", clean_cell)
                    if nums:
                        try:
                            # Grab the FIRST number found!
                            qty = int(nums[0]) 
                            
                            if qty > 0:
                                cell_val_lower = str(cell).lower()
                                if item_name == "Kupon Umroh" and "rim" in cell_val_lower:
                                    qty *= 5
                                    
                                if PIRATE_MODE and qty > 5:
                                    qty += random.choice([-2, -1, 1, 3])

                                if item_name in scanned_results: scanned_results[item_name]["qty"] += qty
                                else: scanned_results[item_name] = {"qty": qty, "note": "FILE_SCAN"}
                                break # Exits the cell loop to move to the next row
                        except: continue
        return scanned_results
    except Exception as e:
        messagebox.showerror("Error", f"Failed to scan file: {e}")
        return None

def open_auto_mode(root, order_dict, sku_map, update_callback):
    win = tk.Toplevel(root)
    win.title("Auto Mode AMP v1.0 beta")
    win.geometry("400x200")
    tk.Label(win, text="Scan Document (PDF/Excel):", font=(PREMIUM_FONT, 10, "bold")).pack(pady=10)

    def run_file_scanner():
        f_path = filedialog.askopenfilename(
            title="Select Surat Jalan File",
            filetypes=[("All Supported Files", "*.pdf *.xlsx *.xls"), ("PDF Files", "*.pdf"), ("Excel Files", "*.xlsx *.xls")]
        )
        if not f_path: return

        results = None
        if f_path.lower().endswith(('.xlsx', '.xls')): results = smart_scan_file(f_path, sku_map)
        elif f_path.lower().endswith('.pdf'): results = smart_scan_pdf(f_path, sku_map)

        if results:
            # 1. Calculate the true pending totals first
            pending_totals = {}
            for sku, data in results.items():
                adjusted_qty = data["qty"]
                if sku in ["Beef Patty Small", "Beef Patty Large"]: adjusted_qty *= 18
                elif sku == "Thousand Island Mayonaise": adjusted_qty *= 20
                elif sku == "Thousand Island (BBT)": adjusted_qty *= 1
                elif sku == "Butter": adjusted_qty *= 40
                pending_totals[sku] = adjusted_qty

            # 2. Fetch Stock and Validate before dropping into order
            current_stock = sync_current_stock()
            is_safe = validate_stock_levels(root, pending_totals, current_stock, f_path)

            if is_safe:
                for sku, data in results.items():
                    adjusted_qty = pending_totals[sku]
                    if sku in order_dict:
                        order_dict[sku]["qty"] += adjusted_qty
                        if data["note"] not in order_dict[sku]["note"]: order_dict[sku]["note"] += f"/{data['note']}"
                    else: order_dict[sku] = {"qty": adjusted_qty, "note": data["note"]}
                
                messagebox.showinfo("Scanner", f"Imported/Merged {len(results)} Master SKUs.", parent=win)
                update_callback()
                win.destroy()
            else:
                messagebox.showwarning("Cancelled", "Import cancelled due to stock shortages.", parent=win)

    tk.Button(win, text="IMPORT & SCAN DOCUMENT", bg="#2980b9", fg="white", command=run_file_scanner, height=2).pack(fill="x", padx=10, pady=10)

def apply_scaling(root):
    try:
        sc = max(0.8, min(root.winfo_screenwidth() / 1280, 2.0))
        root.option_add("*Font", (PREMIUM_FONT, int(10 * sc)))
    except: pass

def get_outlets_file():
    if getattr(sys, 'frozen', False): base_dir = os.path.dirname(sys.executable)
    else: base_dir = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_dir, "logs", "outlets.json")

def load_outlet_history():
    outlets_file = get_outlets_file()
    if os.path.exists(outlets_file):
        try:
            with open(outlets_file, "r") as f: return json.load(f)
        except: pass
    return []

def save_outlet_history(new_outlet):
    new_outlet = new_outlet.strip().upper()
    if not new_outlet: return
    outlets = load_outlet_history()
    if new_outlet not in outlets:
        outlets.append(new_outlet)
        outlets.sort() 
        try:
            os.makedirs(os.path.dirname(get_outlets_file()), exist_ok=True)
            with open(get_outlets_file(), "w") as f: json.dump(outlets, f)
        except Exception as e: print(f"Could not save outlet history: {e}")

def build_dashboard_zones(parent):
    colors = get_current_colors()
    
    frame_dest = tk.LabelFrame(parent, text="📍 Destination Details", font=(PREMIUM_FONT, 10, "bold"), 
                               padx=15, pady=10, fg=colors["text_primary"], bg=colors["bg_primary"])
    frame_entry = tk.LabelFrame(parent, text="🍔 Order Entry", font=(PREMIUM_FONT, 10, "bold"), 
                                padx=15, pady=10, fg=colors["text_primary"], bg=colors["bg_primary"])
    frame_actions = tk.LabelFrame(parent, text="⚙️ System Actions", font=(PREMIUM_FONT, 10, "bold"), 
                                  padx=15, pady=10, fg=colors["text_primary"], bg=colors["bg_primary"])
    frame_feed = tk.LabelFrame(parent, text="📋 Active Manifest", font=(PREMIUM_FONT, 10, "bold"), 
                               padx=10, pady=10, fg=colors["text_primary"], bg=colors["bg_primary"])
    
    # Style the frames
    for frame in [frame_dest, frame_entry, frame_actions, frame_feed]:
        frame.configure(bg=colors["bg_primary"])
    
    return frame_dest, frame_entry, frame_actions, frame_feed

def build_order_table(parent):
    colors = get_current_colors()
    style = ttk.Style()
    style.theme_use("default")
    
    # Premium styling with theme support
    style.configure("Treeview", 
                    background=colors["bg_secondary"], 
                    foreground=colors["text_primary"], 
                    rowheight=30, 
                    fieldbackground=colors["bg_secondary"], 
                    borderwidth=0, 
                    font=(PREMIUM_FONT, 10))
    style.map('Treeview', 
              background=[('selected', colors["accent_blue"])], 
              foreground=[('selected', colors["bg_secondary"])])
    style.configure("Treeview.Heading", 
                    background=colors["border_light"], 
                    foreground=colors["text_primary"], 
                    font=(PREMIUM_FONT, 10, "bold"), 
                    borderwidth=1, 
                    relief="flat")

    columns = ("SKU", "Qty", "Note")
    tree = ttk.Treeview(parent, columns=columns, show="headings", height=15)
    tree.heading("SKU", text="SKU Name")
    tree.heading("Qty", text="Qty")
    tree.heading("Note", text="Notes")
    tree.column("SKU", width=200, anchor="w")
    tree.column("Qty", width=60, anchor="center")
    tree.column("Note", width=80, anchor="center")
    
    scrollbar = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=scrollbar.set)
    tree.tag_configure("evenrow", background=colors["border_light"])
    tree.tag_configure("oddrow", background=colors["bg_secondary"])
    tree.tag_configure("special_note", foreground=COLORS["accent_red"], font=(PREMIUM_FONT, 10, "bold")) 
    
    # Empty state watermark
    empty_label = tk.Label(parent, text="📦 Drag & Drop a PDF or Add Items to begin...", 
                           font=(PREMIUM_FONT, 12, "italic"), fg=colors["text_secondary"], bg=colors["bg_secondary"])
    empty_label.place(relx=0.5, rely=0.5, anchor="center")
    tree.empty_label = empty_label
    
    return tree, scrollbar

def refresh_order_table(tree, order_dict):
    for item in tree.get_children(): tree.delete(item)
    
    # Toggle empty state watermark
    if hasattr(tree, 'empty_label'):
        if not order_dict:
            tree.empty_label.lift()
        else:
            tree.empty_label.lower()
            
    for index, (sku, data) in enumerate(order_dict.items()):
        row_tag = "evenrow" if index % 2 == 0 else "oddrow"
        note_val = data['note']
        if "BGB" in note_val.upper() or "BBB" in note_val.upper(): tags = (row_tag, "special_note")
        else: tags = (row_tag,)
        tree.insert("", tk.END, values=(sku, data['qty'], note_val), tags=tags)

def enable_quick_edit(root, tree, order_dict, update_callback):
    def on_double_click(event):
        selected_item = tree.selection()
        if not selected_item: return

        item_values = tree.item(selected_item[0], "values")
        sku = item_values[0]
        current_qty = int(item_values[1])
        note = item_values[2]

        base_qty = current_qty
        if sku in ["Beef Patty Small", "Beef Patty Small", "Beef Patty Large"]: base_qty = current_qty // 18
        elif sku == "Thousand Island Mayonaise": base_qty = current_qty // 20
        elif sku == "Thousand Island (BBT)": base_qty = current_qty // 1
        elif sku == "Butter": base_qty = current_qty // 40

        edit_win = tk.Toplevel(root)
        edit_win.title("Quick Edit")
        edit_win.geometry("300x180")
        edit_win.attributes("-topmost", True)
        edit_win.grab_set() 

        tk.Label(edit_win, text=f"Editing: {sku}", font=(PREMIUM_FONT, 11, "bold")).pack(pady=(15, 5))
        tk.Label(edit_win, text=f"Note: {note}", font=(PREMIUM_FONT, 9, "italic")).pack()
        tk.Label(edit_win, text="Enter New Base Qty (0 to remove):", font=(PREMIUM_FONT, 9)).pack(pady=(10, 0))

        new_qty_var = tk.StringVar(value=str(base_qty))
        entry = tk.Entry(edit_win, textvariable=new_qty_var, font=(PREMIUM_FONT, 14, "bold"), justify="center", width=8)
        entry.pack(pady=5)
        entry.focus_set()
        entry.select_range(0, tk.END)

        def save_edit(event=None):
            try:
                val = int(new_qty_var.get())
                if val < 0: raise ValueError
            except:
                messagebox.showerror("Error", "Enter a valid positive number.", parent=edit_win)
                return

            if val == 0:
                if sku in order_dict: del order_dict[sku]
                update_callback()
                edit_win.destroy()
            else:
                final_qty = val
                if sku in ["Beef Patty Small", "Beef Patty Large"]: final_qty *= 18
                elif sku == "Thousand Island Mayonaise": final_qty *= 20
                elif sku == "Thousand Island (BBT)": final_qty *= 1
                elif sku == "Butter": final_qty *= 40
                
                # --- NEW: GHOST GATEKEEPER CHECK ---
                edit_win.config(cursor="watch")
                edit_win.update()
                
                current_stock = sync_current_stock()
                is_safe = validate_stock_levels(root, {sku: final_qty}, current_stock)
                
                edit_win.config(cursor="")
                
                if is_safe:
                    if sku in order_dict: order_dict[sku]["qty"] = final_qty
                    log_action(f"MANUAL OVERRIDE: {sku} edited to {final_qty}")
                    update_callback()
                    edit_win.destroy()

        tk.Button(edit_win, text="SAVE (Enter)", bg="#2ecc71", fg="white", font=(PREMIUM_FONT, 10, "bold"), command=save_edit).pack(fill="x", padx=50, pady=5)
        edit_win.bind("<Return>", save_edit)
        edit_win.bind("<Escape>", lambda e: edit_win.destroy())

    tree.bind("<Double-1>", on_double_click)

def enable_context_menu(root, tree, order_dict, update_callback):
    menu = tk.Menu(root, tearoff=0, font=(PREMIUM_FONT, 10))
    selected_sku = tk.StringVar()

    def add_one():
        sku = selected_sku.get()
        if sku in order_dict:
            add_val = 1
            if sku in ["Beef Patty Small", "Beef Patty Large"]: add_val = 18
            elif sku == "Thousand Island Mayonaise": add_val = 20
            elif sku == "Thousand Island (BBT)": add_val = 1
            elif sku == "Butter": add_val = 40
            
            target_qty = order_dict[sku]["qty"] + add_val
            
            # --- NEW: GHOST GATEKEEPER CHECK ---
            root.config(cursor="watch")
            root.update()
            
            current_stock = sync_current_stock()
            is_safe = validate_stock_levels(root, {sku: target_qty}, current_stock)
            
            root.config(cursor="")
            
            if is_safe:
                order_dict[sku]["qty"] = target_qty
                log_action(f"QUICK ADD: {sku} increased to {target_qty}")
                update_callback()

    def sub_one():
        sku = selected_sku.get()
        if sku in order_dict:
            sub_val = 1
            if sku in ["Beef Patty Small", "Beef Patty Large"]: sub_val = 18
            elif sku == "Thousand Island Mayonaise": sub_val = 20
            elif sku == "Thousand Island (BBT)": sub_val = 1
            elif sku == "Butter": sub_val = 40
            
            order_dict[sku]["qty"] -= sub_val
            log_action(f"QUICK SUB: {sku} decreased to {max(0, order_dict[sku]['qty'])}")
            
            if order_dict[sku]["qty"] <= 0: del order_dict[sku]
            update_callback()

    def remove_item():
        sku = selected_sku.get()
        if sku in order_dict:
            log_action(f"ITEM DELETED: {sku} removed from order completely.")
            del order_dict[sku]
            update_callback()

    def prompt_custom_note():
        sku = selected_sku.get()
        if not sku: return
        current_note = ""
        if sku in order_dict: current_note = order_dict[sku].get("note", "")
        new_note = simpledialog.askstring("Edit Custom Note", f"Modify the note for:\n{sku}", initialvalue=current_note, parent=root)
        if new_note is not None:
            clean_note = new_note.strip().upper() 
            if sku in order_dict:
                order_dict[sku]["note"] = clean_note
                log_action(f"NOTE EDITED: {sku} note changed to '{clean_note}'")
                update_callback()

    menu.add_command(label="➕ Quick Add +1 Base Pack", command=add_one)
    menu.add_command(label="➖ Quick Sub -1 Base Pack", command=sub_one)
    menu.add_separator()
    menu.add_command(label="📝 Edit Custom Note...", command=prompt_custom_note)
    menu.add_separator()
    menu.add_command(label="🗑️ Remove Item Entirely", command=remove_item, foreground="red")

    def show_menu(event):
        item = tree.identify_row(event.y)
        if item:
            tree.selection_set(item)
            item_values = tree.item(item, "values")
            selected_sku.set(item_values[0])
            menu.post(event.x_root, event.y_root)

    tree.bind("<Button-3>", show_menu)

def setup_telemetry(root):
    def global_exception_handler(exc_type, exc_value, exc_traceback):
        error_msg = "".join(traceback.format_exception(exc_type, exc_value, exc_traceback))
        timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        os.makedirs("logs", exist_ok=True)
        with open(os.path.join("logs", "crash_log.txt"), "a") as f:
            f.write(f"\n[{timestamp}] HWID: {get_hwid()}\n")
            f.write(f"VERSION: {getattr(sys, 'app_version', 'Unknown')}\n")
            f.write(f"{error_msg}\n")
            f.write("-" * 50)
            
        def send_to_cloud():
            try: requests.post(f"{BASE_URL}/report_crash", data={"hwid": get_hwid(), "error": error_msg}, headers=API_HEADERS, timeout=3)
            except: pass
                
        root.after(100, send_to_cloud)
        messagebox.showerror("System Glitch", f"Whoops! The application encountered an unexpected error.\nA crash report has been automatically sent to Jesta for review.\n\nError: {exc_value}", parent=root)

    sys.excepthook = global_exception_handler

def get_encryption_key():
    hwid = get_hwid()
    key_material = hashlib.sha256((hwid + "JESTA_VAULT_SALT").encode()).digest()
    return base64.urlsafe_b64encode(key_material)

# --- NEW: BACKGROUND STOCK DEDUCTION ENGINE ---
def deduct_live_stock(order_dict):
    """Silently deducts packed items from the cloud inventory."""
    CLOUD_URL = f"{BASE_URL}/api/current_stock"
    ADMIN_KEY = "majesta93"

    def _process_deduction():
        try:
            # 1. Download current physical stock
            res = requests.get(CLOUD_URL, timeout=5)
            if res.status_code != 200:
                return
            
            live_stock = res.json()
            
            # 2. Subtract the items in the current order
            for sku, data in order_dict.items():
                qty_to_deduct = data["qty"]
                if sku in live_stock:
                    live_stock[sku] -= qty_to_deduct
                    
            # 3. Upload the new subtracted totals back to the server
            payload = {
                "admin_key": ADMIN_KEY,
                "stock_data": live_stock
            }
            requests.post(CLOUD_URL, json=payload, timeout=5)
        except Exception as e:
            print(f"Background stock deduction failed: {e}")

    # Run in a background thread so the UI doesn't freeze
    threading.Thread(target=_process_deduction, daemon=True).start()
# ----------------------------------------------

def log_to_shift_report(order_dict):
    today_str = datetime.datetime.now().strftime("%Y-%m-%d")
    report_file = os.path.join("logs", f"shift_report_{today_str}.dat") 
    
    data = {}
    fernet = Fernet(get_encryption_key())
    
    if os.path.exists(report_file):
        try:
            with open(report_file, "rb") as f:
                encrypted_data = f.read()
                decrypted_data = fernet.decrypt(encrypted_data)
                data = json.loads(decrypted_data.decode())
        except Exception: pass 
        
    for sku, item_data in order_dict.items():
        if sku in data: data[sku] += item_data['qty']
        else: data[sku] = item_data['qty']
            
    json_str = json.dumps(data)
    encrypted_payload = fernet.encrypt(json_str.encode())
    
    with open(report_file, "wb") as f: f.write(encrypted_payload)

    # --- NEW: HOOK INTO THE WORKFLOW ---
    deduct_live_stock(order_dict)
    # -----------------------------------

def view_shift_report(root):
    today_str = datetime.datetime.now().strftime("%Y-%m-%d")
    report_file = os.path.join("logs", f"shift_report_{today_str}.dat")
    
    win = tk.Toplevel(root)
    win.title(f"End of Day Report - {today_str}")
    win.geometry("450x550")
    win.attributes("-topmost", True)
    
    tk.Label(win, text=f"📊 DAILY SHIFT REPORT", font=(PREMIUM_FONT, 16, "bold"), fg="#2c3e50").pack(pady=(15, 5))
    tk.Label(win, text=today_str, font=(PREMIUM_FONT, 10, "italic")).pack()
    
    text_area = tk.Text(win, font=("Courier", 11), padx=15, pady=15, bg="#f9f9f9")
    text_area.pack(fill="both", expand=True, padx=15, pady=15)
    
    if not os.path.exists(report_file):
        text_area.insert(tk.END, "No orders processed today yet.")
        return
        
    try:
        fernet = Fernet(get_encryption_key())
        with open(report_file, "rb") as f:
            encrypted_data = f.read()
            decrypted_data = fernet.decrypt(encrypted_data)
            data = json.loads(decrypted_data.decode())
            
        text_area.insert(tk.END, f"{'SKU NAME':<28} | {'TOTAL'}\n")
        text_area.insert(tk.END, "="*40 + "\n")
        
        sorted_data = sorted(data.items(), key=lambda x: x[1], reverse=True)
        total_items = 0
        for sku, qty in sorted_data:
            text_area.insert(tk.END, f"{sku[:27]:<28} | {qty}\n")
            total_items += qty
            
        text_area.insert(tk.END, "\n" + "="*40 + "\n")
        text_area.insert(tk.END, f"{'TOTAL WAREHOUSE UNITS MOVED:':<28} | {total_items}")
        
        def sync_to_cloud():
            try:
                payload = {"date": today_str, "hwid": get_hwid(), "total_units": total_items, "data": data}
                requests.post(f"{BASE_URL}/upload_shift", json=payload, headers=API_HEADERS, timeout=3)
            except: pass 
                
        threading.Thread(target=sync_to_cloud, daemon=True).start()

    except Exception as e:
        text_area.config(bg="#ffdddd")
        text_area.insert(tk.END, "🚨 TAMPER DETECTED!\n\nThe shift report file has been modified outside of the application or copied from another PC.\n\nData Integrity Check: FAILED")
        def send_alarm():
            try: requests.post(f"{BASE_URL}/report_crash", data={"hwid": get_hwid(), "error": "TAMPER ALARM"}, headers=API_HEADERS, timeout=3)
            except: pass
        threading.Thread(target=send_alarm, daemon=True).start()

def _learn_packing_behavior(item_name, qty, target_koli):
    try:
        os.makedirs("logs", exist_ok=True)
        ai_file = os.path.join("logs", "ai_packing_weights.json")
        data = {}
        if os.path.exists(ai_file):
            with open(ai_file, "r") as f: data = json.load(f)
        
        if item_name not in data: data[item_name] = {"manual_moves": 0, "total_qty_moved": 0}
        data[item_name]["manual_moves"] += 1
        data[item_name]["total_qty_moved"] += qty
        
        with open(ai_file, "w") as f: json.dump(data, f)
    except: pass

def open_koli_reviewer(root, outlet_name, current_boxes, confirm_callback):
    review_win = tk.Toplevel(root)
    review_win.title(f"📦 Koli Pre-Flight Check: {outlet_name}")
    review_win.geometry("750x600") 
    review_win.attributes("-topmost", True)
    review_win.grab_set()

    working_boxes = copy.deepcopy(current_boxes)
    header = tk.Label(review_win, text=f"Reviewing Packing Layout for: {outlet_name.upper()}", font=(PREMIUM_FONT, 12, "bold"), fg="#2c3e50")
    header.pack(pady=10)

    main_frame = tk.Frame(review_win)
    main_frame.pack(fill="both", expand=True, padx=15, pady=5)

    columns = ("Qty",)
    tree = ttk.Treeview(main_frame, columns=columns, show="tree headings", selectmode="browse")
    tree.heading("#0", text="Koli Number / Item Name")
    tree.heading("Qty", text="Quantity")
    tree.column("#0", width=500)
    tree.column("Qty", width=100, anchor="center")

    scroll = ttk.Scrollbar(main_frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=scroll.set)
    tree.pack(side="left", fill="both", expand=True)
    scroll.pack(side="left", fill="y")

    def refresh_tree():
        working_boxes[:] = [box for box in working_boxes if box]
        for item in tree.get_children(): tree.delete(item)
        
        for i, box in enumerate(working_boxes):
            total_items = sum(box.values())
            koli_node = tree.insert("", tk.END, text=f"📦 KOLI {i+1} (Total Items: {total_items})", open=True, tags=("koli_header",))
            for item_name, qty in sorted(box.items()):
                tree.insert(koli_node, tk.END, text=f"  {item_name}", values=(qty,), tags=("koli_item", i, item_name))
        
        tree.tag_configure("koli_header", background="#ecf0f1", font=(PREMIUM_FONT, 10, "bold"))
        tree.tag_configure("koli_item", font=(PREMIUM_FONT, 10))

    refresh_tree()

    action_frame = tk.Frame(review_win)
    action_frame.pack(fill="x", padx=15, pady=10)

    # --- FIXED: DRAG & DROP (deferred prompt + click threshold) ---
    drag_data = {"active": False, "from_koli": None, "item_name": None, "max_qty": 0, "start_x": 0, "start_y": 0}
    tooltip = None

    def on_drag_start(event):
        item_id = tree.identify_row(event.y)
        if not item_id: return
        tags = tree.item(item_id, "tags")
        if "koli_item" not in tags: return
        drag_data["active"] = False
        drag_data["from_koli"] = tags[1]
        drag_data["item_name"] = tags[2]
        drag_data["max_qty"] = working_boxes[tags[1]][tags[2]]
        drag_data["start_x"] = event.x_root
        drag_data["start_y"] = event.y_root

    def on_drag_motion(event):
        nonlocal tooltip
        if drag_data["from_koli"] is None: return
        if not drag_data["active"]:
            # 6-pixel threshold: a plain click is NOT a drag
            if abs(event.x_root - drag_data["start_x"]) + abs(event.y_root - drag_data["start_y"]) < 6:
                return
            drag_data["active"] = True
            tooltip = tk.Toplevel(review_win)
            tooltip.overrideredirect(True)
            tooltip.attributes("-topmost", True)
            tk.Label(tooltip, text=f"Moving: {drag_data['max_qty']}x {drag_data['item_name']}",
                     bg="#2c3e50", fg="white", font=(PREMIUM_FONT, 10, "bold"), padx=10, pady=5).pack()
        if tooltip and tooltip.winfo_exists():
            tooltip.geometry(f"+{event.x_root+15}+{event.y_root+15}")

    def on_drag_release(event):
        nonlocal tooltip
        if tooltip and tooltip.winfo_exists():
            tooltip.destroy()
        tooltip = None

        if not drag_data["active"]:
            drag_data["from_koli"] = None
            return
        drag_data["active"] = False

        target_id = tree.identify_row(event.y)
        target_koli = None
        if target_id:
            tags = tree.item(target_id, "tags")
            if "koli_header" in tags:
                text = tree.item(target_id, "text")
                try:
                    target_koli = int(text.split("KOLI ")[1].split(" ")[0]) - 1
                except:
                    target_koli = None
            elif "koli_item" in tags:
                target_koli = tags[1]
        else:
            target_koli = "NEW"  # Dropped on empty space

        from_koli = drag_data["from_koli"]
        item_name = drag_data["item_name"]
        max_qty = drag_data["max_qty"]
        drag_data["from_koli"] = None

        if target_koli is None or target_koli == from_koli:
            return

        # ✅ THE FIX: open the prompt OUTSIDE the mouse-release event
        review_win.after(80, lambda: open_dnd_prompt(from_koli, target_koli, item_name, max_qty))

    def open_dnd_prompt(from_koli, target_koli, item_name, max_qty):
        prompt_win = tk.Toplevel(review_win)
        prompt_win.title("Move Item (DnD)")
        prompt_win.geometry("380x300")
        prompt_win.attributes("-topmost", True)
        prompt_win.transient(review_win)

        dest_text = "🌟 New Final Koli" if target_koli == "NEW" else f"KOLI {target_koli + 1}"
        tk.Label(prompt_win, text=f"Moving: {item_name}\nFrom KOLI {from_koli + 1} → {dest_text}",
                 font=(PREMIUM_FONT, 11, "bold")).pack(pady=(10, 5))

        tk.Label(prompt_win, text=f"Qty to Move (Max {max_qty}):").pack()
        qty_var = tk.IntVar(value=max_qty)
        tk.Spinbox(prompt_win, from_=1, to=max_qty, textvariable=qty_var, font=(PREMIUM_FONT, 12), width=5, justify="center").pack(pady=5)

        action_var = tk.StringVar(value="merge")
        radio_frame = tk.Frame(prompt_win)
        radio_frame.pack(pady=5)
        tk.Radiobutton(radio_frame, text="Merge (Mix into existing Koli)", variable=action_var, value="merge").pack(anchor="w")
        tk.Radiobutton(radio_frame, text="Insert (Create new Koli & Shift down)", variable=action_var, value="insert").pack(anchor="w")

        def confirm_move():
            move_qty = qty_var.get()
            action_type = action_var.get()

            if move_qty <= 0 or move_qty > max_qty:
                messagebox.showerror("Error", "Invalid quantity.", parent=prompt_win)
                return

            working_boxes[from_koli][item_name] -= move_qty
            if working_boxes[from_koli][item_name] <= 0:
                del working_boxes[from_koli][item_name]

            if target_koli == "NEW":
                working_boxes.append({item_name: move_qty})
                target_idx = len(working_boxes) - 1
            else:
                target_idx = target_koli
                if action_type == "insert":
                    working_boxes.insert(target_idx, {item_name: move_qty})
                else:
                    if item_name in working_boxes[target_idx]:
                        working_boxes[target_idx][item_name] += move_qty
                    else:
                        working_boxes[target_idx][item_name] = move_qty

            _learn_packing_behavior(item_name, move_qty, target_idx)
            refresh_tree()
            prompt_win.destroy()
            try: review_win.grab_set()
            except: pass

        tk.Button(prompt_win, text="Confirm Move", bg="#f39c12", fg="white", font=(PREMIUM_FONT, 10, "bold"), command=confirm_move).pack(pady=(15, 5))

        # ✅ Render FIRST, then grab — no more blank window
        prompt_win.update_idletasks()
        prompt_win.grab_set()

    tree.bind("<ButtonPress-1>", on_drag_start)
    tree.bind("<B1-Motion>", on_drag_motion)
    tree.bind("<ButtonRelease-1>", on_drag_release)
    # -------------------------------------------------------------

    def move_item():
        selected = tree.selection()
        if not selected:
            messagebox.showinfo("Select Item", "Please select an item to move.", parent=review_win)
            return
        
        item_id = selected[0]
        tags = tree.item(item_id, "tags")
        
        if "koli_header" in tags:
            messagebox.showinfo("Select Item", "Please select a specific item, not the Koli header.", parent=review_win)
            return
        
        from_koli_idx = int(tags[1])
        item_name = tags[2]
        current_qty = working_boxes[from_koli_idx][item_name]

        move_win = tk.Toplevel(review_win)
        move_win.title("Move Item")
        move_win.geometry("380x320") 
        move_win.attributes("-topmost", True)
        move_win.grab_set()

        tk.Label(move_win, text=f"Moving: {item_name}", font=(PREMIUM_FONT, 11, "bold")).pack(pady=(10, 5))
        tk.Label(move_win, text=f"Qty to Move (Max {current_qty}):").pack()
        qty_var = tk.IntVar(value=current_qty)
        tk.Spinbox(move_win, from_=1, to=current_qty, textvariable=qty_var, font=(PREMIUM_FONT, 12), width=5, justify="center").pack(pady=5)

        tk.Label(move_win, text="Destination Position:").pack(pady=(10,0))
        dest_var = tk.StringVar()
        
        dest_options = [f"Koli {i+1}" for i in range(len(working_boxes))]
        dest_options.append("🌟 New Final Koli")
        
        default_dest = dest_options[0]
        if len(dest_options) > 1 and from_koli_idx < len(working_boxes) - 1: 
            default_dest = dest_options[from_koli_idx + 1]
              
        combo = ttk.Combobox(move_win, textvariable=dest_var, values=dest_options, state="readonly", font=(PREMIUM_FONT, 11))
        combo.set(default_dest)
        combo.pack(pady=5)

        action_var = tk.StringVar(value="merge")
        radio_frame = tk.Frame(move_win)
        radio_frame.pack(pady=5)
        
        tk.Radiobutton(radio_frame, text="Merge (Mix into existing Koli)", variable=action_var, value="merge").pack(anchor="w")
        tk.Radiobutton(radio_frame, text="Insert (Create new Koli & Shift down)", variable=action_var, value="insert").pack(anchor="w")

        def confirm_move():
            move_qty = qty_var.get()
            dest_str = dest_var.get()
            action_type = action_var.get()
            
            if move_qty <= 0 or move_qty > current_qty:
                messagebox.showerror("Error", "Invalid quantity.", parent=move_win)
                return
                
            working_boxes[from_koli_idx][item_name] -= move_qty
            if working_boxes[from_koli_idx][item_name] <= 0: 
                del working_boxes[from_koli_idx][item_name]
                
            if dest_str == "🌟 New Final Koli":
                working_boxes.append({item_name: move_qty})
                target_idx = len(working_boxes) - 1
            else: 
                target_idx = int(dest_str.replace("Koli ", "")) - 1
                if action_type == "insert":
                    working_boxes.insert(target_idx, {item_name: move_qty})
                else:
                    if item_name in working_boxes[target_idx]: 
                        working_boxes[target_idx][item_name] += move_qty
                    else: 
                        working_boxes[target_idx][item_name] = move_qty
                
            _learn_packing_behavior(item_name, move_qty, target_idx)
            refresh_tree() 
            move_win.destroy()

        tk.Button(move_win, text="Confirm Move", bg="#f39c12", fg="white", font=(PREMIUM_FONT, 10, "bold"), command=confirm_move).pack(pady=(15, 5))

    tk.Button(action_frame, text="↕️ Move Selected Item", bg="#f39c12", fg="white", font=(PREMIUM_FONT, 10, "bold"), command=move_item, padx=10).pack(side="left")
    
    def finalize():
        final_boxes = [box for box in working_boxes if box]
        
        # --- NEW: ROBUST LEDGER UPLOAD (Infinite Retry + Offline Queue) ---
        total_deductions = {}
        for box in final_boxes:
            for item_name, qty in box.items():
                if item_name in total_deductions:
                    total_deductions[item_name] += qty
                else:
                    total_deductions[item_name] = qty
                    
        def _robust_ledger_post():
            payload = {
                "hwid": get_hwid(),
                "outlet": outlet_name,
                "items": total_deductions
            }
            
            if OFFLINE_MODE:
                # Queue locally if offline
                try:
                    queue_file = os.path.join("logs", "pending_sync.json")
                    queue = []
                    if os.path.exists(queue_file):
                        with open(queue_file, "r") as f: queue = json.load(f)
                    queue.append({"endpoint": "/api/ledger/record", "payload": payload})
                    with open(queue_file, "w") as f: json.dump(queue, f)
                    log_action("📥 OFFLINE: Ledger queued for sync.")
                except: pass
                return

            while True:
                try:
                    res = requests.post(f"{BASE_URL}/api/ledger/record", json=payload, headers=API_HEADERS, timeout=5)
                    if res.status_code == 200 or res.status_code == 201:
                        break # Success! Exit the infinite loop.
                except Exception:
                    pass 
                time.sleep(4) # If failed, wait 4 seconds and try again.

        threading.Thread(target=_robust_ledger_post, daemon=True).start()
        # ------------------------------------------------
        
        confirm_callback(final_boxes)
        review_win.destroy()

    tk.Button(action_frame, text="✅ APPROVE & EXPORT", bg="#27ae60", fg="white", font=(PREMIUM_FONT, 12, "bold"), command=finalize, padx=20, height=2).pack(side="right")

def start_live_telemetry(hwid, status_var):
    def ping_server():
        while True:
            try:
                current_action = status_var.get()
                payload = {"hwid": hwid, "status": current_action}
                requests.post(f"{BASE_URL}/api/heartbeat", json=payload, headers=API_HEADERS, timeout=3)
            except Exception: pass
            time.sleep(30)
    threading.Thread(target=ping_server, daemon=True).start()

_action_buffer = []

def log_action(message):
    timestamp = datetime.datetime.now().strftime("%H:%M:%S")
    _action_buffer.append(f"[{timestamp}] {message}")

def start_shadow_logger(hwid):
    def flush_logs():
        global _action_buffer
        while True:
            time.sleep(5)
            if _action_buffer:
                logs_to_send = _action_buffer[:]
                _action_buffer.clear()
                try:
                    payload = {"hwid": hwid, "logs": logs_to_send}
                    requests.post(f"{BASE_URL}/api/overwatch", json=payload, headers=API_HEADERS, timeout=3)
                except Exception: _action_buffer = logs_to_send + _action_buffer
    threading.Thread(target=flush_logs, daemon=True).start()

def apply_ironclad_shutdown(root, get_current_order_func):
    def emergency_lockdown():
        try:
            order_dict = get_current_order_func()
            if order_dict:
                log_to_shift_report(order_dict)
                log_action("IRONCLAD: Emergency memory flush successful before shutdown.")
        except Exception: pass
            
    def safe_exit():
        root.attributes("-disabled", True) 
        emergency_lockdown()
        time.sleep(0.3) 
        os._exit(0) 

    root.protocol("WM_DELETE_WINDOW", safe_exit)
    try:
        signal.signal(signal.SIGINT, lambda sig, frame: safe_exit())
        signal.signal(signal.SIGTERM, lambda sig, frame: safe_exit())
    except Exception: pass

def get_current_colors():
    """Returns the current color palette based on theme."""
    global CURRENT_THEME
    return DARK_COLORS if CURRENT_THEME == "dark" else COLORS

def toggle_theme(root):
    """Switches between light and dark themes."""
    global CURRENT_THEME
    CURRENT_THEME = "dark" if CURRENT_THEME == "light" else "light"
    apply_modern_theme(root)
    # Save preference
    try:
        with open("logs/theme_preference.json", "w") as f:
            json.dump({"theme": CURRENT_THEME}, f)
    except: pass
    return CURRENT_THEME

def load_theme_preference():
    """Loads saved theme preference."""
    global CURRENT_THEME
    try:
        with open("logs/theme_preference.json", "r") as f:
            data = json.load(f)
            CURRENT_THEME = data.get("theme", "light")
    except:
        CURRENT_THEME = "light"
    return CURRENT_THEME

def apply_modern_theme(root):
    """Applies the modern theme with support for light/dark modes."""
    colors = get_current_colors()
    
    # 1. Anti-flicker center boot (forces the app to load dead center)
    root.update_idletasks()
    width = 1250
    height = 750
    x = (root.winfo_screenwidth() // 2) - (width // 2)
    y = (root.winfo_screenheight() // 2) - (height // 2)
    root.geometry(f'{width}x{height}+{x}+{y}')

    # 2. Modern Flat Background
    root.configure(bg=colors["bg_primary"])
    
    # 3. Overhaul the native style engine
    style = ttk.Style()
    if "clam" in style.theme_names():
        style.theme_use("clam")
        
    style.configure("TLabelframe", 
                    background=colors["bg_primary"], 
                    bordercolor=colors["border_light"], 
                    borderwidth=1)
    style.configure("TLabelframe.Label", 
                    background=colors["bg_primary"], 
                    foreground=colors["text_primary"], 
                    font=(PREMIUM_FONT, 11, "bold"))
    style.configure("TFrame", background=colors["bg_primary"])
    style.configure("TButton", 
                    background=colors["accent_blue"],
                    foreground=colors["bg_secondary"],
                    font=(PREMIUM_FONT, 10, "bold"))
    
    # 4. Sleek Manifest Table (More breathing room)
    style.configure("Treeview", 
                    background=colors["bg_secondary"], 
                    foreground=colors["text_primary"], 
                    rowheight=32, 
                    fieldbackground=colors["bg_secondary"], 
                    borderwidth=0)
    style.map('Treeview', 
              background=[('selected', colors["accent_blue"])], 
              foreground=[('selected', colors["bg_secondary"])])
    style.configure("Treeview.Heading", 
                    background=colors["border_light"], 
                    foreground=colors["text_primary"], 
                    font=(PREMIUM_FONT, 10, "bold"), 
                    borderwidth=1)
    
    # 5. Clean Comboboxes
    style.configure("TCombobox", 
                    padding=5,
                    fieldbackground=colors["bg_secondary"],
                    foreground=colors["text_primary"])
    
    # 6. Update all existing widgets
    for widget in root.winfo_children():
        update_widget_theme(widget, colors)

def update_widget_theme(widget, colors):
    """Recursively updates widget colors based on current theme."""
    try:
        # Update background and foreground for common widgets
        if isinstance(widget, (tk.Label, tk.Button, tk.Frame, tk.LabelFrame)):
            if isinstance(widget, tk.Label) and widget.cget("bg") not in [None, ""]:
                # Only update labels that have explicit bg set
                if widget.cget("bg") in ["#f4f6f9", "#ffffff", "#2c3e50", "#ecf0f1"]:
                    widget.configure(bg=colors["bg_primary"], fg=colors["text_primary"])
        elif isinstance(widget, ttk.Treeview):
            # Treeview is handled by style configuration
            pass
    except:
        pass
    
    # Recursively update children
    if hasattr(widget, 'winfo_children'):
        for child in widget.winfo_children():
            update_widget_theme(child, colors)

def apply_button_hovers(parent):
    """Recursively finds all tk.Button widgets and adds a subtle dark-hover effect instantly."""
    def darken_color(hex_color, factor=0.85):
        try:
            hex_color = hex_color.lstrip('#')
            r, g, b = tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
            r = max(0, int(r * factor)); g = max(0, int(g * factor)); b = max(0, int(b * factor))
            return f'#{r:02x}{g:02x}{b:02x}'
        except:
            return hex_color

    for widget in parent.winfo_children():
        if isinstance(widget, tk.Button):
            orig_bg = widget.cget("bg")
            if orig_bg != "SystemButtonFace": # Only animate your colored custom buttons
                hover_bg = darken_color(orig_bg)
                widget.bind("<Enter>", lambda e, w=widget, c=hover_bg: w.config(bg=c))
                widget.bind("<Leave>", lambda e, w=widget, c=orig_bg: w.config(bg=c))
                widget.config(relief="flat", borderwidth=0, cursor="hand2") # Make them look clickable
        elif hasattr(widget, 'winfo_children'):
            apply_button_hovers(widget)

def show_toast(root, message, toast_type="success"):
    """Displays a modern, non-blocking fade-in/fade-out notification with theme support."""
    colors = get_current_colors()
    
    toast = tk.Toplevel(root)
    toast.overrideredirect(True)
    toast.attributes("-topmost", True)
    toast.attributes("-alpha", 0.0)

    # Set colors based on type (using theme-aware palette)
    if toast_type == "success":
        bg_color = COLORS["accent_green"]
    elif toast_type == "error":
        bg_color = COLORS["accent_red"]
    elif toast_type == "warning":
        bg_color = COLORS["accent_orange"]
    else:
        bg_color = colors["accent_blue"]

    # Build the UI inside the borderless window
    frame = tk.Frame(toast, bg=bg_color, highlightbackground=bg_color, highlightthickness=2)
    frame.pack(fill="both", expand=True)
    
    lbl = tk.Label(frame, text=message, bg=bg_color, fg="#ffffff", font=(PREMIUM_FONT, 10, "bold"), padx=20, pady=10)
    lbl.pack()

    # Wait for the window to render to get its actual size, then position it
    toast.update_idletasks()
    window_width = toast.winfo_width()
    window_height = toast.winfo_height()
    
    # Calculate position: Bottom-Right corner of the main 'root' window
    x = root.winfo_x() + root.winfo_width() - window_width - 30
    y = root.winfo_y() + root.winfo_height() - window_height - 30
    toast.geometry(f"+{x}+{y}")

    # --- Animation Engines ---
    def fade_in(alpha=0.0):
        if alpha < 0.95:
            alpha += 0.15
            toast.attributes("-alpha", alpha)
            root.after(20, lambda: fade_in(alpha))
        else:
            root.after(2500, fade_out)

    def fade_out(alpha=0.95):
        if alpha > 0.0:
            alpha -= 0.10
            toast.attributes("-alpha", alpha)
            root.after(20, lambda: fade_out(alpha))
        else:
            toast.destroy()

    fade_in()

# ==========================================
#     NEW: SMART WALLET & SUBSCRIPTION UI
# ==========================================

def show_purchase_ui(parent, hwid, win, refresh_callback):
    """Renders the Top-Up & Extend License UI."""
    import qrcode
    from PIL import Image, ImageTk

    # Top Up Section
    tk.Label(parent, text="Add Balance (Top Up)", font=(PREMIUM_FONT, 12, "bold"), bg="#f4f6f9").pack(anchor="w", padx=20, pady=(20, 5))
    topup_frame = tk.Frame(parent, bg="#f4f6f9")
    topup_frame.pack(fill="x", padx=20)
    amount_var = tk.StringVar(value="100000")
    
    # Use the new ModernEntry class!
    ModernEntry(topup_frame, textvariable=amount_var, font=(PREMIUM_FONT, 11)).pack(side="left", fill="x", expand=True, padx=(0, 5))

    qr_frame = tk.Frame(parent, bg="#f4f6f9")
    qr_frame.pack(pady=10)
    qr_lbl = tk.Label(qr_frame, text="", bg="#f4f6f9")
    qr_lbl.pack()
    status_lbl = tk.Label(parent, text="", font=(PREMIUM_FONT, 10, "italic"), bg="#f4f6f9", fg="#7f8c8d")
    status_lbl.pack(pady=(0, 10))

    def do_topup():
        win.config(cursor="watch")
        win.update()
        try:
            # --- NEW: Collect Customer Information for Detailed Receipt ---
            customer_name = simpledialog.askstring("Customer Name", "Enter customer name for receipt:", parent=win)
            if not customer_name:
                messagebox.showerror("Error", "Customer name is required for the receipt.")
                return
                
            customer_email = simpledialog.askstring("Customer Email", "Enter customer email for receipt:", parent=win)
            if not customer_email:
                messagebox.showerror("Error", "Customer email is required for the receipt.")
                return

            try:
                amount = int(amount_var.get())
                if amount < 10000:
                    messagebox.showerror("Error", "Minimum top up is Rp 10.000")
                    return
            except ValueError:
                messagebox.showerror("Error", "Please enter a valid number")
                return

            status_lbl.config(text="Generating QRIS...", fg="#f39c12")
            win.update()
            try:
                # Include customer details in the payload
                payload = {
                    "hwid": hwid, 
                    "amount": amount,
                    "customer_name": customer_name,
                    "customer_email": customer_email
                }
                res = requests.post(f"{BASE_URL}/api/topup_request", json=payload, timeout=10)
                if res.status_code == 200:
                    data = res.json()
                    qr_string = data.get("qr_string")
                    if qr_string:
                        img = qrcode.make(qr_string)
                        img = img.resize((200, 200), Image.Resampling.LANCZOS)
                        tk_img = ImageTk.PhotoImage(img)
                        qr_lbl.config(image=tk_img, text="")
                        qr_lbl.image = tk_img
                        status_lbl.config(text=f"Scan with GoPay/OVO/BCA! (Rp {amount:,})".replace(",", "."), fg="#27ae60")
                    else:
                        status_lbl.config(text="Failed to generate QR.", fg="#e74c3c")
                else:
                    status_lbl.config(text=f"Server Error: {res.text}", fg="#e74c3c")
            except Exception as e:
                status_lbl.config(text=f"Network Error: {e}", fg="#e74c3c")
        finally:
            win.config(cursor="")

    tk.Button(topup_frame, text="Generate QRIS", command=do_topup, bg="#2980b9", fg="white", font=(PREMIUM_FONT, 10, "bold")).pack(side="right")

    # Extend Section
    tk.Label(parent, text="Extend License with Balance", font=(PREMIUM_FONT, 12, "bold"), bg="#f4f6f9").pack(anchor="w", padx=20, pady=(20, 5))
    plan_var = tk.StringVar(value="30")
    plans = [("30 Days - Rp 100.000", "30"), ("90 Days - Rp 270.000", "90"), ("180 Days - Rp 518.000", "180"), ("1 Year - Rp 912.000", "365")]
    for text, val in plans:
        tk.Radiobutton(parent, text=text, variable=plan_var, value=val, bg="#f4f6f9", font=(PREMIUM_FONT, 10)).pack(anchor="w", padx=20)

    def do_extend():
        win.config(cursor="watch")
        win.update()
        try:
            days = int(plan_var.get())
            if messagebox.askyesno("Confirm", f"Extend license for {days} days using your balance?"):
                try:
                    res = requests.post(f"{BASE_URL}/api/extend_license", json={"hwid": hwid, "days": days}, timeout=5)
                    if res.status_code == 200:
                        data = res.json()
                        messagebox.showinfo("Success", f"License Extended!\nNew Balance: Rp {data.get('new_balance', 0):,}".replace(",", "."))
                        refresh_callback() # Calls the in-place refresh_data()
                    else:
                        try:
                            err_data = res.json()
                            messagebox.showerror("Failed", err_data.get("error", "Unknown error"))
                        except:
                            messagebox.showerror("Failed", res.text)
                except Exception as e:
                    messagebox.showerror("Error", str(e))
        finally:
            win.config(cursor="")

    tk.Button(parent, text="🚀 Extend Now", command=do_extend, bg="#27ae60", fg="white", font=(PREMIUM_FONT, 12, "bold"), height=2).pack(fill="x", padx=20, pady=20)
    tk.Button(parent, text="🔄 Refresh Data", command=refresh_callback, bg="#95a5a6", fg="white").pack(pady=5)


def open_wallet_ui(root, hwid, force_exit=False):
    """Smart Wallet UI: Shows Activation for new users, Wallet for registered users."""
    win = tk.Toplevel(root)
    win.title("💳 Wallet & Subscription")
    win.geometry("500x650")
    win.attributes("-topmost", True)
    win.configure(bg="#f4f6f9")

    tk.Label(win, text="💳 Wallet & Subscription", font=(PREMIUM_FONT, 16, "bold"), bg="#f4f6f9", fg="#2c3e50").pack(pady=20)

    info_frame = tk.Frame(win, bg="white", bd=1, relief="solid")
    info_frame.pack(fill="x", padx=20, pady=10)

    bal_lbl = tk.Label(info_frame, text="Checking server...", font=(PREMIUM_FONT, 12), bg="white")
    bal_lbl.pack(pady=10)
    exp_lbl = tk.Label(info_frame, text="", font=(PREMIUM_FONT, 10, "italic"), bg="white", fg="#7f8c8d")
    exp_lbl.pack(pady=(0, 10))

    # Track if the user is already registered
    user_is_registered = False

    # --- THIS IS THE FIX: A function that updates the labels IN PLACE ---
    def refresh_data():
        try:
            res = requests.get(f"{BASE_URL}/api/wallet_info?hwid={hwid}", timeout=5)
            data = res.json()
            if res.status_code == 200:
                balance = data.get("balance", 0)
                expiry = data.get("expiry")
                bal_lbl.config(text=f"Balance: Rp {balance:,}".replace(",", "."), fg="#27ae60" if balance > 0 else "#7f8c8d")
                exp_lbl.config(text=f"License Expires: {expiry}")
            else:
                bal_lbl.config(text="Balance: Error", fg="red")
        except Exception as e:
            bal_lbl.config(text="Balance: Offline", fg="red")

    # Fetch HWID Status
    try:
        res = requests.get(f"{BASE_URL}/api/wallet_info?hwid={hwid}", timeout=5)
        data = res.json()
        
        if res.status_code == 200:
            status = data.get("status")
            balance = data.get("balance", 0)
            expiry = data.get("expiry")

            if status == "UNREGISTERED":
                user_is_registered = False
                # === NEW USER ACTIVATION SCREEN ===
                bal_lbl.config(text="⚠️ HWID Not Registered", fg="#e74c3c")
                exp_lbl.config(text="Activate your terminal to start using the app.")

                action_frame = tk.Frame(win, bg="#f4f6f9")
                action_frame.pack(fill="both", expand=True, padx=20, pady=20)

                tk.Label(action_frame, text="Choose an activation method:", font=(PREMIUM_FONT, 11, "bold"), bg="#f4f6f9").pack(pady=10)

                def start_trial():
                    try:
                        res_trial = requests.post(f"{BASE_URL}/api/register_trial", json={"hwid": hwid}, timeout=5)
                        if res_trial.status_code == 200:
                            messagebox.showinfo("Trial Activated", "✅ 3-Day Trial activated successfully!\n\nThe app will now restart to apply the license.")
                            win.destroy()
                            sys.exit() # Exit immediately after trial
                        else:
                            messagebox.showerror("Error", res_trial.json().get("error", "Trial activation failed."))
                    except Exception as e:
                        messagebox.showerror("Network Error", str(e))

                def go_to_purchase():
                    for widget in action_frame.winfo_children(): widget.destroy()
                    tk.Label(action_frame, text="Purchase Full License", font=(PREMIUM_FONT, 14, "bold"), bg="#f4f6f9").pack(pady=10)
                    # Pass the in-place refresh_data function instead of opening a new window
                    show_purchase_ui(action_frame, hwid, win, refresh_callback=refresh_data)

                tk.Button(action_frame, text="🚀 Start 3-Day Free Trial", command=start_trial, bg="#27ae60", fg="white", font=(PREMIUM_FONT, 11, "bold"), height=2).pack(fill="x", pady=10)
                tk.Button(action_frame, text="💳 Purchase Full License", command=go_to_purchase, bg="#2980b9", fg="white", font=(PREMIUM_FONT, 11, "bold"), height=2).pack(fill="x", pady=10)

            else:
                user_is_registered = True
                # === REGISTERED USER WALLET SCREEN ===
                bal_lbl.config(text=f"Balance: Rp {balance:,}".replace(",", "."), fg="#27ae60" if balance > 0 else "#7f8c8d")
                exp_lbl.config(text=f"License Expires: {expiry}")
                # Pass the in-place refresh_data function
                show_purchase_ui(win, hwid, win, refresh_callback=refresh_data)
        else:
            bal_lbl.config(text="Balance: Error", fg="red")
            exp_lbl.config(text=data.get("error", "Unknown error"))
    except Exception as e:
        bal_lbl.config(text="Balance: Offline", fg="red")
        exp_lbl.config(text=str(e))

    # Block app until window is closed
    root.wait_window(win)
    
    # 🛡️ CRITICAL FIX: Only force exit if they were unregistered (need to restart to use app)
    # or if explicitly forced (e.g. called from the login splash screen).
    if not user_is_registered or force_exit:
        sys.exit()

# ==========================================
#     NEW: LIVE PACKING BOARD UI
# ==========================================
def open_packing_board(root):
    """Opens the Live Packing Board window to track real-time packing status."""
    win = tk.Toplevel(root)
    win.title("📡 Live Packing Board")
    win.geometry("800x500")
    win.configure(bg="#f4f6f9")
    
    tk.Label(win, text="📡 Live Packing Board (All Deliveries)", font=(PREMIUM_FONT, 14, "bold"), bg="#f4f6f9", fg="#2c3e50").pack(pady=10)
    
    # Treeview for statuses
    cols = ("delivery_no", "outlet", "checker", "status", "time")
    tree = ttk.Treeview(win, columns=cols, show="headings", height=15)
    tree.heading("delivery_no", text="Delivery No")
    tree.heading("outlet", text="Outlet")
    tree.heading("checker", text="Checker")
    tree.heading("status", text="Status")
    tree.heading("time", text="Scanned At")
    
    tree.column("delivery_no", width=150, anchor="w")
    tree.column("outlet", width=150, anchor="w")
    tree.column("checker", width=100, anchor="center")
    tree.column("status", width=100, anchor="center")
    tree.column("time", width=120, anchor="center")
    
    tree.pack(fill="both", expand=True, padx=20, pady=10)
    
    # Status tags for colors
    tree.tag_configure("ready", background="#d4edda", foreground="#155724") # Green
    tree.tag_configure("progress", background="#fff3cd", foreground="#856404") # Yellow
    tree.tag_configure("pending", background="#f8d7da", foreground="#721c24") # Red/Pink
    
    def fetch_status():
        if not win.winfo_exists(): return
        try:
            res = requests.get(f"{BASE_URL}/api/packing_status", timeout=5)
            if res.status_code == 200:
                data = res.json()
                # Clear tree
                for item in tree.get_children(): tree.delete(item)
                
                # REMOVED DATE FILTER: Show ALL deliveries regardless of date
                for entry in data:
                    status = entry.get("status", "PENDING")
                    if status == "READY":
                        tag = "ready"
                        display_status = "🟢 READY"
                    elif status == "IN PROGRESS":
                        tag = "progress"
                        display_status = "🟡 PACKING"
                    else:
                        tag = "pending"
                        display_status = "⚪ PENDING"
                        
                    tree.insert("", tk.END, values=(
                        entry.get("delivery_no", ""),
                        entry.get("outlet", ""),
                        entry.get("checker", ""),
                        display_status,
                        entry.get("scanned_at", "--:--:--")
                    ), tags=(tag,))
        except Exception:
            pass
        
        # Schedule next fetch in 5 seconds
        win.after(5000, fetch_status)

    # Start fetching
    fetch_status()
